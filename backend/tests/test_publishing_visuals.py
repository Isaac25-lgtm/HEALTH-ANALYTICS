"""Downloads carry the dashboard's colour coding, labels and district map, from snapshot values only."""

from uuid import UUID

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from sqlalchemy import select

from app.config import get_settings
from app.domain.enums import AggregationClass
from app.models import ExportJob, Geometry, OrgUnit
from app.services.export_view import STATUS_FILL, format_measure, resolve_status, status_label
from app.services.geography import descendants, top_units_of_class
from app.services.publishing_pptx import SCORECARD_TABLE_NAME, UNITS_TABLE_NAME
from tests.conftest import auth_header, login, query_dashboard
from tests.helpers import put_population, put_raw
from tests.test_amendment_exports import _export


def _unit(session, code):
    return session.scalar(select(OrgUnit).where(OrgUnit.code == code))


def _square(lon: float, lat: float) -> dict:
    return {
        "type": "Polygon",
        "coordinates": [[[lon, lat], [lon + 0.4, lat], [lon + 0.4, lat + 0.4], [lon, lat + 0.4], [lon, lat]]],
    }


def _districts(session):
    uganda = _unit(session, "UG")
    return sorted(
        top_units_of_class(descendants(session, uganda), AggregationClass.DISTRICT_EQUIVALENT), key=lambda u: u.code
    )


def _seed(session):
    # A national value needs the complete district cohort, so every fixture district reports.
    for index, unit in enumerate(_districts(session)):
        put_population(session, unit, 2024, 100_000, code=f"VIS_POP_{unit.code}")
        put_raw(session, unit, "FY2024/25", "ANC1", 4_000 + 500 * index)
        put_raw(session, unit, "FY2024/25", "ANC4", 2_000)
        session.add(Geometry(org_unit_id=unit.id, geojson=_square(32 + index, 2 + index), geometry_kind="polygon"))
    put_population(session, _unit(session, "UG"), 2024, 300_000, code="VIS_POP_UG")
    session.commit()


def _national(client, session, tmp_path, monkeypatch):
    _seed(session)
    monkeypatch.setattr(get_settings(), "export_dir", str(tmp_path))
    headers = auth_header(login(client, "national.analyst"))
    uganda = _unit(session, "UG")
    return headers, uganda, query_dashboard(client, headers, uganda.id).json()


def _path(session, job) -> str:
    return session.get(ExportJob, UUID(job["job_id"])).file_path


def test_excel_scorecard_and_unit_grid_are_colour_coded(client, session, tmp_path, monkeypatch):
    headers, uganda, dash = _national(client, session, tmp_path, monkeypatch)
    book = load_workbook(_path(session, _export(client, headers, uganda, dash, "excel")))
    assert {"Scorecard", "Units", "Ranking", "Trends"} <= set(book.sheetnames)
    score = book["Scorecard"]
    rows = {row[0].value: row for row in score.iter_rows(min_row=2) if row[0].value}
    anc1 = next(item for item in dash["module_result"]["indicators"] if item["indicator_code"] == "ANC1_COVERAGE")
    status = resolve_status(anc1)
    value_cell, status_cell = rows[anc1["name"]][1], rows[anc1["name"]][3]
    assert value_cell.value == anc1["raw_value"]
    assert value_cell.fill.fgColor.rgb.endswith(STATUS_FILL[status])
    assert status_cell.value == status_label(status)
    assert value_cell.number_format.endswith('"%"')
    units = book["Units"]
    names = [row[0].value for row in units.iter_rows(min_row=6)]
    districts = {row["org_unit_name"] for row in dash["module_result"]["district_comparison"]}
    assert districts <= set(names)


def test_powerpoint_is_widescreen_with_map_charts_and_coloured_tables(client, session, tmp_path, monkeypatch):
    headers, uganda, dash = _national(client, session, tmp_path, monkeypatch)
    deck = Presentation(_path(session, _export(client, headers, uganda, dash, "powerpoint")))
    assert round(deck.slide_width / deck.slide_height, 2) == 1.78
    freeforms, charts, pictures, table_names = 0, 0, 0, set()
    coloured_cells = 0
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.FREEFORM:
                freeforms += 1
            if shape.has_chart:
                charts += 1
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if shape.has_table:
                table_names.add(shape.name)
                for row in list(shape.table.rows)[1:]:
                    for cell in row.cells:
                        if cell.fill.type is not None and str(cell.fill.fore_color.rgb) in STATUS_FILL.values():
                            coloured_cells += 1
    assert freeforms == len(_districts(session))  # one vector shape per district with a boundary
    assert charts >= 1 and pictures == 0
    assert {SCORECARD_TABLE_NAME, UNITS_TABLE_NAME} <= table_names
    assert coloured_cells > 0


def test_pdf_and_word_reports_are_real_documents_with_status_colours(client, session, tmp_path, monkeypatch):
    headers, uganda, dash = _national(client, session, tmp_path, monkeypatch)
    pdf_path = _path(session, _export(client, headers, uganda, dash, "pdf"))
    with open(pdf_path, "rb") as handle:
        assert handle.read(5) == b"%PDF-"
    document = Document(_path(session, _export(client, headers, uganda, dash, "word")))
    text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    anc1 = next(item for item in dash["module_result"]["indicators"] if item["indicator_code"] == "ANC1_COVERAGE")
    assert anc1["name"] in text
    assert format_measure(anc1) in text
    shaded = document.element.xpath("//w:shd/@w:fill")
    assert any(fill in STATUS_FILL.values() for fill in shaded)


def test_blob_transport_returns_the_same_bytes_as_raw_octets(client, session, tmp_path, monkeypatch):
    headers, uganda, dash = _national(client, session, tmp_path, monkeypatch)
    job = _export(client, headers, uganda, dash, "pdf")
    governed = client.post(f"/exports/jobs/{job['job_id']}/download", headers=headers)
    blob = client.post(f"/exports/jobs/{job['job_id']}/download?transport=blob", headers=headers)
    assert governed.headers["content-type"] == "application/pdf"
    assert blob.headers["content-type"] == "application/octet-stream"
    assert "content-disposition" not in blob.headers
    assert blob.headers["x-export-filename"].endswith(".pdf")
    assert blob.content == governed.content and blob.content.startswith(b"%PDF-")
    rejected = client.post(f"/exports/jobs/{job['job_id']}/download?transport=inline", headers=headers)
    assert rejected.status_code == 422
