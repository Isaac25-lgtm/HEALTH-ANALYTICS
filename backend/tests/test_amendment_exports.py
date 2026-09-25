"""Amendment §15: exports are complete, never silently truncated, and accurately labelled."""

from pathlib import Path
from uuid import UUID

from openpyxl import load_workbook
from pptx import Presentation
from sqlalchemy import select

from app.config import get_settings
from app.domain.exports import EXPORT_FORMATS
from app.domain.modules import MODULE_INDICATORS
from app.models import ExportJob, OrgUnit
from app.services.publishing import NOT_RECORDED
from app.services.publishing_pptx import SCORECARD_TABLE_NAME
from tests.conftest import auth_header, download_export, login, query_dashboard
from tests.helpers import put_population, put_raw


def _unit(session, code):
    return session.scalar(select(OrgUnit).where(OrgUnit.code == code))


def _export(client, headers, org, dash, kind):
    created = client.post(
        f"/exports/{kind}",
        json={
            "org_unit_id": str(org.id),
            "period": dash["period"],
            "module": dash["module"],
            "analysis_snapshot_id": dash["analysis_snapshot_id"],
            "view_hash": dash["view_hash"],
        },
        headers=headers,
    )
    assert created.status_code == 202, created.text
    return created.json()


def test_excel_governance_columns_are_never_blank(client, session, tmp_path, monkeypatch):
    pader = _unit(session, "PADER")
    put_population(session, pader, 2024, 1_000_000, code="EXPORT_GOV_POP")
    put_raw(session, pader, "FY2024/25", "ANC1", 47_700)
    put_raw(session, pader, "FY2024/25", "ANC4", 29_650)
    session.commit()
    monkeypatch.setattr(get_settings(), "export_dir", str(tmp_path))
    headers = auth_header(login(client, "pader.focal"))
    dash = query_dashboard(client, headers, pader.id).json()
    job = _export(client, headers, pader, dash, "excel")
    book = load_workbook(tmp_path / f"{job['job_id']}.xlsx")
    calc = list(book["Calculations"].iter_rows(values_only=True))
    raw = list(book["Raw Data"].iter_rows(values_only=True))
    calc_header, raw_header = calc[0], raw[0]
    governed_calc = [calc_header.index(name) for name in ("Formula version", "Thresholds", "Change interpretation")]
    governed_raw = [raw_header.index(name) for name in ("Mapping version", "Source freshness", "Aggregation scope")]
    assert len(calc) - 1 == len(MODULE_INDICATORS["anc"])
    for row in calc[1:]:
        assert all(row[index] not in (None, "") for index in governed_calc), row
    for row in raw[1:]:
        assert all(row[index] not in (None, "") for index in governed_raw), row
    by_code = {row[0]: row for row in calc[1:]}
    assert "Green >=95%" in by_code["ANC1_COVERAGE"][calc_header.index("Thresholds")]
    raw_by_code = {row[0]: row for row in raw[1:]}
    assert raw_by_code["ANC1_COVERAGE"][raw_header.index("Mapping version")] == "v1"
    assert raw_by_code["ANC1_COVERAGE"][raw_header.index("Source freshness")] != NOT_RECORDED


def test_powerpoint_includes_every_indicator_of_a_large_module(client, session, tmp_path, monkeypatch):
    uganda = _unit(session, "UG")
    put_population(session, uganda, 2024, 10_000_000, code="EXPORT_EPI_POP")
    put_raw(session, uganda, "FY2024/25", "BCG", 50_000)
    session.commit()
    monkeypatch.setattr(get_settings(), "export_dir", str(tmp_path))
    headers = auth_header(login(client, "national.analyst"))
    dash = query_dashboard(client, headers, uganda.id, module="immunization").json()
    job = _export(client, headers, uganda, dash, "powerpoint")
    deck = Presentation(Path(session.get(ExportJob, UUID(job["job_id"])).file_path))
    table_rows: list[str] = []
    notes = ""
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_table and shape.name == SCORECARD_TABLE_NAME:
                table_rows.extend(row.cells[0].text for index, row in enumerate(shape.table.rows) if index > 0)
            elif shape.has_text_frame and "indicators in the snapshot" in shape.text_frame.text:
                notes = shape.text_frame.text
    expected = [row["name"] for row in dash["module_result"]["indicators"]]
    assert len(expected) == len(MODULE_INDICATORS["immunization"]) == 30
    assert table_rows == expected
    assert "All 30 indicators" in notes
    assert "Immunization briefing" in deck.slides[0].shapes.title.text


def test_labels_media_types_and_extensions_are_accurate(client, session, tmp_path, monkeypatch):
    pader = _unit(session, "PADER")
    put_population(session, pader, 2024, 100_000, code="EXPORT_LABEL_POP")
    put_raw(session, pader, "FY2024/25", "ANC1", 100)
    put_raw(session, pader, "FY2024/25", "ANC4", 50)
    session.commit()
    monkeypatch.setattr(get_settings(), "export_dir", str(tmp_path))
    headers = auth_header(login(client, "pader.focal"))
    dash = query_dashboard(client, headers, pader.id).json()
    actions = {item["kind"]: item for item in dash["exports"]["actions"]}
    for kind, spec in EXPORT_FORMATS.items():
        assert actions[kind]["label"] == spec["label"]
        assert actions[kind]["format"] == spec["format"]
        job = _export(client, headers, pader, dash, kind)
        assert spec["label"] in job["message"]
        meta = client.get(f"/exports/jobs/{job['job_id']}", headers=headers).json()
        assert meta["label"] == spec["label"]
        download = download_export(client, headers, job["job_id"])
        assert download.status_code == 200
        assert download.headers["content-type"].startswith(spec["media_type"].split(";")[0])
        assert download.headers["content-disposition"].endswith(f'{spec["extension"]}"')
    assert "Markdown" in actions["report"]["label"]
    assert "Word" not in actions["report"]["label"] and "PDF" not in actions["report"]["label"]
    for kind in ("word", "pdf"):
        assert actions[kind]["available"] is True
        assert actions[kind]["implemented"] is True
