"""Retention windows, the purge service and purge-safe provenance.

Covers work packages C (configurable retention), D (retained evidence survives raw deletion)
and the export-file half of F (expired artifacts are explainable, not 404s).
"""

from __future__ import annotations

import json
from datetime import UTC, datetime, timedelta
from uuid import UUID, uuid4

import pytest
from sqlalchemy import func, select, update

from app.config import Settings, get_settings, validate_runtime_settings
from app.domain import retention
from app.domain.retention import policy_summary, retention_cutoffs, subtract_months
from app.models import (
    AnalysisSnapshot,
    AuditLog,
    CalculatedValue,
    CalculationRun,
    DataQualityFlag,
    ExportArtifact,
    ExportJob,
    MaintenanceLock,
    MaintenanceRun,
    OrgUnit,
    RawAggregateValue,
    RawEventSnapshot,
)
from app.services import artifact_store, purge
from app.services.purge import (
    STATUS_COMPLETED,
    STATUS_DISABLED,
    STATUS_FAILED,
    STATUS_LOCKED,
    purge_expired,
    purge_policy,
)
from tests.conftest import auth_header, download_export, login, query_dashboard
from tests.helpers import put_population, put_raw

OLD = datetime(2020, 1, 1, tzinfo=UTC)


def _unit(session, code="PADER"):
    return session.scalar(select(OrgUnit).where(OrgUnit.code == code))


def _count(session, model) -> int:
    return int(session.scalar(select(func.count()).select_from(model)) or 0)


# ---------------------------------------------------------------------------
# Cutoffs and configuration
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("moment", "months", "expected"),
    [
        (datetime(2026, 9, 14, tzinfo=UTC), 36, datetime(2023, 9, 14, tzinfo=UTC)),
        (datetime(2026, 9, 14, tzinfo=UTC), 24, datetime(2024, 9, 14, tzinfo=UTC)),
        (datetime(2026, 3, 31, tzinfo=UTC), 1, datetime(2026, 2, 28, tzinfo=UTC)),
        (datetime(2024, 2, 29, tzinfo=UTC), 12, datetime(2023, 2, 28, tzinfo=UTC)),
        (datetime(2026, 1, 15, tzinfo=UTC), 13, datetime(2024, 12, 15, tzinfo=UTC)),
    ],
)
def test_month_windows_use_calendar_months(moment, months, expected):
    assert subtract_months(moment, months) == expected


def test_default_windows_match_the_owner_decision():
    settings = Settings(_env_file=None)
    assert settings.raw_aggregate_retention_days == 7
    assert settings.mpdsr_event_retention_hours == 24
    assert settings.export_file_retention_hours == 24
    assert settings.export_job_retention_days == 90
    assert settings.calculation_snapshot_retention_months == 36
    assert settings.audit_log_retention_months == 24
    now = datetime(2026, 9, 14, 12, 0, tzinfo=UTC)
    cutoffs = retention_cutoffs(settings, now)
    assert cutoffs[retention.RAW_AGGREGATES] == now - timedelta(days=7)
    assert cutoffs[retention.MPDSR_EVENTS] == now - timedelta(hours=24)
    assert cutoffs[retention.SNAPSHOTS] == datetime(2023, 9, 14, 12, 0, tzinfo=UTC)
    assert cutoffs[retention.AUDIT_LOGS] == datetime(2024, 9, 14, 12, 0, tzinfo=UTC)
    assert {item["policy"] for item in policy_summary(settings)} == set(retention.POLICY_ORDER)


def test_production_rejects_incoherent_or_unsafe_retention_settings():
    base = {
        "_env_file": None,
        "app_env": "production",
        "database_url": "postgresql+psycopg://hpip:owner-supplied@db:5432/hpip",
        "auth_secret": "a-long-production-secret-value-that-is-not-a-placeholder",
        "seed_password": "not-the-development-default",
        "auth_cookie_secure": True,
        "web_origin": "https://hpip.example.test",
        "allowed_origins": "https://hpip.example.test",
        "rate_limit_backend": "redis",
        "redis_url": "redis://:owner@redis:6379/0",
        "celery_broker_url": "redis://:owner@redis:6379/1",
        "celery_result_backend": "redis://:owner@redis:6379/2",
        # The test environment enables eager exports; production settings must not inherit it.
        "export_eager": False,
    }
    assert validate_runtime_settings(Settings(**base)) == []
    assert any(
        "EXPORT_JOB_RETENTION_DAYS" in error
        for error in validate_runtime_settings(
            Settings(**base, export_job_retention_days=1, export_file_retention_hours=48)
        )
    )
    assert any("PURGE_ENABLED" in error for error in validate_runtime_settings(Settings(**base, purge_enabled=False)))
    assert any(
        "PURGE_DRY_RUN" in error
        for error in validate_runtime_settings(Settings(**base, purge_dry_run=True, purge_schedule_enabled=True))
    )
    assert any(
        "EXPORT_ARTIFACT_STORAGE" in error
        for error in validate_runtime_settings(Settings(**base, export_shared_filesystem=False))
    )
    assert (
        validate_runtime_settings(Settings(**base, export_artifact_storage="database", export_shared_filesystem=False))
        == []
    )


# ---------------------------------------------------------------------------
# Raw caches
# ---------------------------------------------------------------------------


def _raw_rows(session, org, *, fresh: int, expired: int) -> None:
    now = datetime.now(UTC)
    for index in range(fresh):
        put_raw(session, org, "FY2024/25", "ANC1", 100 + index)
    for row in session.scalars(select(RawAggregateValue)).all():
        row.extracted_at = now
    for index in range(expired):
        put_raw(session, org, "FY2023/24", "ANC4", 200 + index)
    session.flush()
    for row in session.scalars(select(RawAggregateValue).where(RawAggregateValue.period == "FY2023/24")).all():
        row.extracted_at = now - timedelta(days=30)
    session.commit()


def test_expired_raw_aggregates_are_deleted_and_fresh_rows_are_kept(session):
    org = _unit(session)
    _raw_rows(session, org, fresh=2, expired=3)
    result = purge_policy(session, retention.RAW_AGGREGATES, dry_run=False)
    assert result.status == STATUS_COMPLETED
    assert result.rows_examined == 3 and result.rows_deleted == 3
    remaining = session.scalars(select(RawAggregateValue)).all()
    assert len(remaining) == 2
    assert {row.period for row in remaining} == {"FY2024/25"}
    # Idempotent: a second pass has nothing left to do.
    again = purge_policy(session, retention.RAW_AGGREGATES, dry_run=False)
    assert again.rows_examined == 0 and again.rows_deleted == 0


def test_dry_run_changes_nothing_at_all(session):
    org = _unit(session)
    _raw_rows(session, org, fresh=1, expired=2)
    before_rows = _count(session, RawAggregateValue)
    before_runs = _count(session, MaintenanceRun)
    before_locks = _count(session, MaintenanceLock)
    results = purge_expired(session, dry_run=True)
    assert all(item.dry_run for item in results)
    raw = next(item for item in results if item.policy == retention.RAW_AGGREGATES)
    assert raw.rows_examined == 2 and raw.rows_deleted == 0
    assert _count(session, RawAggregateValue) == before_rows
    assert _count(session, MaintenanceRun) == before_runs
    assert _count(session, MaintenanceLock) == before_locks


def test_expired_mpdsr_events_and_their_uids_disappear_everywhere(session):
    org = _unit(session)
    now = datetime.now(UTC)
    fresh = RawEventSnapshot(
        event_uid="EVENTFRESH1",
        org_unit_id=org.id,
        status="COMPLETED",
        extracted_at=now,
        data_values={"cause_category": "Haemorrhage"},
    )
    stale = RawEventSnapshot(
        event_uid="EVENTSTALE1",
        org_unit_id=org.id,
        status="COMPLETED",
        extracted_at=now - timedelta(hours=48),
        data_values={"cause_category": "Sepsis"},
    )
    session.add_all([fresh, stale])
    session.flush()
    flag = DataQualityFlag(
        fingerprint="retention-flag-1",
        rule_id="MPDSR_INCOMPLETE_DATES",
        severity="WARNING",
        status="open",
        org_unit_id=org.id,
        event_uid="EVENTSTALE1",
        explanation="Synthetic retention flag",
        first_detected_at=now - timedelta(hours=48),
        last_detected_at=now - timedelta(hours=48),
    )
    session.add(flag)
    session.commit()

    result = purge_policy(session, retention.MPDSR_EVENTS, dry_run=False)
    assert result.rows_deleted == 1
    remaining = session.scalars(select(RawEventSnapshot)).all()
    assert [row.event_uid for row in remaining] == ["EVENTFRESH1"]
    session.refresh(flag)
    # The flag survives as an aggregate data-quality record, without the temporary event UID.
    assert flag.event_uid is None
    assert flag.rule_id == "MPDSR_INCOMPLETE_DATES"


# ---------------------------------------------------------------------------
# Exports
# ---------------------------------------------------------------------------


def _export(client, session, headers, org, tmp_path, monkeypatch, kind="excel"):
    monkeypatch.setattr(get_settings(), "export_dir", str(tmp_path))
    dash = query_dashboard(client, headers, org.id).json()
    created = client.post(
        f"/exports/{kind}",
        json={
            "org_unit_id": str(org.id),
            "period": "FY2024/25",
            "module": "anc",
            "analysis_snapshot_id": dash["analysis_snapshot_id"],
            "view_hash": dash["view_hash"],
        },
        headers=headers,
    )
    assert created.status_code == 202, created.text
    return dash, created.json()["job_id"]


def test_expired_export_file_is_explained_while_job_metadata_survives(client, session, tmp_path, monkeypatch):
    org = _unit(session)
    put_population(session, org, 2024, 1_000_000, code="RET_POP")
    put_raw(session, org, "FY2024/25", "ANC1", 40_000)
    session.commit()
    headers = auth_header(login(client, "pader.focal"))
    _dash, job_id = _export(client, session, headers, org, tmp_path, monkeypatch)
    job = session.get(ExportJob, UUID(job_id))
    checksum, size = job.checksum, job.artifact_size_bytes
    assert job.artifact_expires_at is not None
    assert download_export(client, headers, job_id).status_code == 200

    # Age the artifact past its window.
    session.execute(
        update(ExportJob)
        .where(ExportJob.id == job.id)
        .values(artifact_expires_at=datetime.now(UTC) - timedelta(hours=1))
    )
    session.commit()
    result = purge_policy(session, retention.EXPORT_FILES, dry_run=False)
    assert result.files_examined == 1 and result.files_deleted == 1

    session.expire_all()
    job = session.get(ExportJob, UUID(job_id))
    assert job.status == "succeeded"
    assert job.artifact_deleted_at is not None
    assert job.checksum == checksum and job.artifact_size_bytes == size
    assert not list(tmp_path.glob("*.xlsx"))

    expired = download_export(client, headers, job_id)
    assert expired.status_code == 410
    detail = expired.json()["detail"]
    assert detail["code"] == artifact_store.ARTIFACT_EXPIRED
    assert detail["checksum"] == checksum
    meta = client.get(f"/exports/jobs/{job_id}", headers=headers).json()
    assert meta["downloadable"] is False and meta["artifact_expired"] is True
    # A repeat request regenerates from the same snapshot rather than serving nothing.
    again = client.post(
        "/exports/excel",
        json={
            "org_unit_id": str(org.id),
            "period": "FY2024/25",
            "module": "anc",
            "analysis_snapshot_id": str(job.analysis_snapshot_id),
            "view_hash": job.view_hash,
        },
        headers=headers,
    )
    assert again.status_code == 202
    assert again.json()["job_id"] == job_id


def test_database_artifact_storage_round_trips_without_a_shared_filesystem(client, session, tmp_path, monkeypatch):
    org = _unit(session)
    put_population(session, org, 2024, 1_000_000, code="RET_DB_POP")
    put_raw(session, org, "FY2024/25", "ANC1", 40_000)
    session.commit()
    monkeypatch.setattr(get_settings(), "export_artifact_storage", "database")
    headers = auth_header(login(client, "pader.focal"))
    _dash, job_id = _export(client, session, headers, org, tmp_path, monkeypatch)
    job = session.get(ExportJob, UUID(job_id))
    assert job.artifact_storage == "database"
    assert job.file_path is None
    assert not list(tmp_path.iterdir())  # nothing left on the local disk
    stored = session.scalar(select(ExportArtifact).where(ExportArtifact.export_job_id == job.id))
    assert stored is not None and stored.size_bytes == job.artifact_size_bytes
    download = download_export(client, headers, job_id)
    assert download.status_code == 200
    assert len(download.content) == stored.size_bytes
    assert download.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    session.execute(
        update(ExportJob)
        .where(ExportJob.id == job.id)
        .values(artifact_expires_at=datetime.now(UTC) - timedelta(hours=1))
    )
    session.execute(
        update(ExportArtifact)
        .where(ExportArtifact.export_job_id == job.id)
        .values(expires_at=datetime.now(UTC) - timedelta(hours=1))
    )
    session.commit()
    purge_policy(session, retention.EXPORT_FILES, dry_run=False)
    assert session.scalar(select(ExportArtifact).where(ExportArtifact.export_job_id == job.id)) is None
    assert download_export(client, headers, job_id).status_code == 410


def test_terminal_export_metadata_expires_but_live_jobs_never_do(session):
    org = _unit(session)
    admin = session.scalar(select(OrgUnit).where(OrgUnit.code == "PADER"))
    del admin
    user_id = session.scalars(select(ExportJob.user_id)).first()
    if user_id is None:
        from app.models import User

        user_id = session.scalar(select(User.id).where(User.username == "pader.focal"))
    old = datetime.now(UTC) - timedelta(days=400)
    jobs = [
        ExportJob(
            id=uuid4(), user_id=user_id, export_type="excel", status="succeeded", finished_at=old, org_unit_id=org.id
        ),
        ExportJob(
            id=uuid4(), user_id=user_id, export_type="excel", status="failed", finished_at=old, org_unit_id=org.id
        ),
        ExportJob(id=uuid4(), user_id=user_id, export_type="excel", status="queued", org_unit_id=org.id),
        ExportJob(
            id=uuid4(), user_id=user_id, export_type="excel", status="running", claimed_at=old, org_unit_id=org.id
        ),
    ]
    for job in jobs:
        job.created_at = old
    session.add_all(jobs)
    session.commit()
    result = purge_policy(session, retention.EXPORT_JOBS, dry_run=False)
    assert result.rows_deleted == 2
    statuses = {row.status for row in session.scalars(select(ExportJob)).all()}
    assert statuses == {"queued", "running"}


# ---------------------------------------------------------------------------
# Snapshots and audit
# ---------------------------------------------------------------------------


def test_expired_snapshots_take_their_runs_and_values_but_spare_referenced_ones(client, session, tmp_path, monkeypatch):
    org = _unit(session)
    put_population(session, org, 2024, 1_000_000, code="RET_SNAP_POP")
    put_raw(session, org, "FY2024/25", "ANC1", 40_000)
    session.commit()
    headers = auth_header(login(client, "pader.focal"))
    keep_dash, keep_job = _export(client, session, headers, org, tmp_path, monkeypatch)
    drop_dash = query_dashboard(client, headers, org.id, request_key=str(uuid4())).json()
    old = datetime.now(UTC) - timedelta(days=40 * 30)
    for snapshot_id in (keep_dash["analysis_snapshot_id"], drop_dash["analysis_snapshot_id"]):
        session.execute(update(AnalysisSnapshot).where(AnalysisSnapshot.id == UUID(snapshot_id)).values(created_at=old))
    session.execute(update(CalculationRun).values(created_at=old))
    session.commit()
    runs_before = _count(session, CalculationRun)
    values_before = _count(session, CalculatedValue)

    result = purge_policy(session, retention.SNAPSHOTS, dry_run=False)
    assert result.rows_skipped >= 1
    remaining = {str(row.id) for row in session.scalars(select(AnalysisSnapshot)).all()}
    assert keep_dash["analysis_snapshot_id"] in remaining  # still referenced by a retained export job
    assert drop_dash["analysis_snapshot_id"] not in remaining
    assert _count(session, CalculationRun) < runs_before
    assert _count(session, CalculatedValue) < values_before
    # The retained snapshot still reopens and its export job still resolves.
    reopened = client.get(f"/analysis-snapshots/{keep_dash['analysis_snapshot_id']}", headers=headers)
    assert reopened.status_code == 200
    assert client.get(f"/exports/jobs/{keep_job}", headers=headers).status_code == 200


def test_audit_log_retention(session):
    old = AuditLog(
        action="legacy_action", resource_type="export_job", created_at=datetime.now(UTC) - timedelta(days=800)
    )
    recent = AuditLog(action="recent_action", resource_type="export_job", created_at=datetime.now(UTC))
    session.add_all([old, recent])
    session.commit()
    result = purge_policy(session, retention.AUDIT_LOGS, dry_run=False)
    assert result.rows_deleted >= 1
    actions = {row.action for row in session.scalars(select(AuditLog)).all()}
    assert "legacy_action" not in actions
    assert "recent_action" in actions


# ---------------------------------------------------------------------------
# Locking, failure handling and the run record
# ---------------------------------------------------------------------------


def test_second_purge_is_skipped_while_the_lease_is_held(session, engine):
    from sqlalchemy.orm import sessionmaker

    other = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)()
    try:
        assert purge.acquire_lease(other, "retention_purge:audit_logs", "other-holder", 600)
        result = purge_policy(session, retention.AUDIT_LOGS, dry_run=False)
        assert result.status == STATUS_LOCKED
        run = session.scalars(select(MaintenanceRun).order_by(MaintenanceRun.started_at.desc())).first()
        assert run is None or run.status != STATUS_COMPLETED
        purge.release_lease(other, "retention_purge:audit_logs", "other-holder")
        # Once released, the same policy runs normally and the lease is not left behind.
        assert purge_policy(session, retention.AUDIT_LOGS, dry_run=False).status == STATUS_COMPLETED
        assert session.get(MaintenanceLock, "retention_purge:audit_logs") is None
    finally:
        other.close()


def test_expired_lease_can_be_taken_over(session):
    stale = datetime.now(UTC) - timedelta(hours=2)
    session.add(MaintenanceLock(name="retention_purge:test", holder="dead-worker", acquired_at=stale, expires_at=stale))
    session.commit()
    assert purge.acquire_lease(session, "retention_purge:test", "new-worker", 600) is True
    assert session.get(MaintenanceLock, "retention_purge:test").holder == "new-worker"


def test_failure_is_recorded_safely_and_releases_the_lease(session, monkeypatch):
    def explode(*_args, **_kwargs):
        raise RuntimeError("database connection lost at /var/data/hpip.db")

    monkeypatch.setitem(purge.HANDLERS, retention.AUDIT_LOGS, explode)
    result = purge_policy(session, retention.AUDIT_LOGS, dry_run=False)
    assert result.status == STATUS_FAILED
    assert result.error_code == "purge_failed"
    assert result.error_summary == "RuntimeError"
    assert session.get(MaintenanceLock, "retention_purge:audit_logs") is None
    run = session.scalars(
        select(MaintenanceRun)
        .where(MaintenanceRun.policy == retention.AUDIT_LOGS)
        .order_by(MaintenanceRun.started_at.desc())
    ).first()
    assert run is not None and run.status == STATUS_FAILED
    encoded = json.dumps({column.name: str(getattr(run, column.name)) for column in MaintenanceRun.__table__.columns})
    assert "/var/data" not in encoded and "connection lost" not in encoded


def test_run_records_capture_counts_without_content(session):
    org = _unit(session)
    _raw_rows(session, org, fresh=1, expired=1)
    result = purge_policy(session, retention.RAW_AGGREGATES, dry_run=False, source="scheduler")
    run = session.get(MaintenanceRun, UUID(result.run_id))
    assert run.task_type == "retention_purge"
    assert run.entity == "raw_aggregate_values"
    assert run.source == "scheduler"
    assert run.rows_deleted == 1 and run.batches >= 1
    assert run.requested_cutoff is not None and run.software_version
    assert run.error_code is None


def test_purge_can_be_disabled(session, monkeypatch):
    monkeypatch.setattr(get_settings(), "purge_enabled", False)
    results = purge_expired(session)
    assert {item.status for item in results} == {STATUS_DISABLED}
    assert _count(session, MaintenanceRun) == 0


def test_batching_is_bounded(session, monkeypatch):
    org = _unit(session)
    _raw_rows(session, org, fresh=0, expired=5)
    monkeypatch.setattr(get_settings(), "purge_batch_size", 2)
    monkeypatch.setattr(get_settings(), "purge_max_batches", 2)
    result = purge_policy(session, retention.RAW_AGGREGATES, dry_run=False)
    assert result.batches == 2 and result.rows_deleted == 4
    assert _count(session, RawAggregateValue) == 1  # the rest waits for the next scheduled run


def test_unknown_policy_is_rejected(session):
    with pytest.raises(ValueError, match="Unknown retention policy"):
        purge_expired(session, policy_names=["not_a_policy"])


# ---------------------------------------------------------------------------
# Work package D: retained evidence survives raw deletion
# ---------------------------------------------------------------------------


def test_snapshot_and_export_survive_the_deletion_of_their_raw_rows(client, session, tmp_path, monkeypatch):
    from openpyxl import load_workbook

    org = _unit(session)
    put_population(session, org, 2024, 1_000_000, code="PROV_POP")
    put_raw(session, org, "FY2024/25", "ANC1", 47_700)
    put_raw(session, org, "FY2024/25", "ANC4", 29_650)
    session.commit()
    headers = auth_header(login(client, "pader.focal"))
    dash, job_id = _export(client, session, headers, org, tmp_path, monkeypatch)
    before = {
        row["indicator_code"]: (row["raw_value"], row["numerator"], row["denominator"], row["display_value"])
        for row in dash["module_result"]["indicators"]
    }
    run_id = dash["module_result"]["current_run_id"]
    snapshot_id = dash["analysis_snapshot_id"]
    first_download = download_export(client, headers, job_id)
    assert first_download.status_code == 200

    # Expire and purge every raw row the run used.
    session.execute(update(RawAggregateValue).values(extracted_at=datetime.now(UTC) - timedelta(days=30)))
    session.commit()
    raw_result = purge_policy(session, retention.RAW_AGGREGATES, dry_run=False)
    assert raw_result.rows_deleted >= 2
    assert _count(session, RawAggregateValue) == 0

    # 1. The committed snapshot reopens with identical values.
    reopened = client.get(f"/analysis-snapshots/{snapshot_id}", headers=headers)
    assert reopened.status_code == 200
    after = {
        row["indicator_code"]: (row["raw_value"], row["numerator"], row["denominator"], row["display_value"])
        for row in reopened.json()["module_result"]["indicators"]
    }
    assert after == before

    # 2. Provenance still explains the published result without the raw rows.
    stored_run = session.get(CalculationRun, UUID(run_id))
    lineage = stored_run.config_snapshot
    assert lineage["software_version"] and lineage["indicator_versions"]
    assert lineage["raw_row_ids"]  # historical identifiers are retained
    versions = {item["indicator_code"]: item for item in lineage["indicator_versions"]}
    assert versions["ANC1_COVERAGE"]["formula_version"]
    assert versions["ANC1_COVERAGE"]["formula_version_rule"]
    from app.models import Indicator, IndicatorVersion

    value = session.scalar(
        select(CalculatedValue)
        .join(IndicatorVersion, IndicatorVersion.id == CalculatedValue.indicator_version_id)
        .join(Indicator, Indicator.id == IndicatorVersion.indicator_id)
        .where(
            CalculatedValue.calculation_run_id == stored_run.id,
            CalculatedValue.org_unit_id == org.id,
            Indicator.code == "ANC1_COVERAGE",
        )
    )
    assert value.numerator is not None and value.denominator is not None
    assert float(value.numerator) == 47_700
    assert value.population_year and value.software_version and value.aggregation_policy
    assert value.source_row_ids  # the identifiers remain even though the rows are gone

    # 3. A fresh export regenerates from the snapshot with the same numbers.
    regenerated = client.post(
        "/exports/excel",
        json={
            "org_unit_id": str(org.id),
            "period": "FY2024/25",
            "module": "anc",
            "analysis_snapshot_id": snapshot_id,
            "view_hash": dash["view_hash"],
            "comparison_period": dash.get("comparison_period"),
        },
        headers=headers,
    )
    assert regenerated.status_code == 202, regenerated.text
    new_job = regenerated.json()["job_id"]
    book = load_workbook(tmp_path / f"{new_job}.xlsx")
    rows = {row[0]: row for row in book["Calculations"].iter_rows(min_row=2, values_only=True)}
    assert rows["ANC1_COVERAGE"][2] == before["ANC1_COVERAGE"][0]
    assert rows["ANC4_COVERAGE"][2] == before["ANC4_COVERAGE"][0]


def test_event_uids_are_gone_after_purge_but_aggregate_output_remains(client, session, monkeypatch):
    org = _unit(session)
    now = datetime.now(UTC)
    session.add(
        RawEventSnapshot(
            event_uid="EVENTPURGE01",
            org_unit_id=org.id,
            status="COMPLETED",
            extracted_at=now - timedelta(days=2),
            data_values={"cause_category": "Sepsis"},
        )
    )
    session.commit()
    purge_policy(session, retention.MPDSR_EVENTS, dry_run=False)
    assert _count(session, RawEventSnapshot) == 0
    headers = auth_header(login(client, "admin.user"))
    response = query_dashboard(client, headers, org.id, module="mpdsr")
    assert response.status_code == 201, response.text
    encoded = json.dumps(response.json())
    assert "EVENTPURGE01" not in encoded
    assert response.json()["module_result"]["module"] == "mpdsr"
