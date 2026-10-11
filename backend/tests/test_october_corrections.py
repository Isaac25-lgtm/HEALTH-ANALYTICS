"""Independent regressions for the authorised correction pass; fixtures never reach DHIS2."""

from datetime import UTC, date, datetime
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import func, select, update

from app.models import AnalyticsRevision, CalculationRun, EventFieldMapping, Geometry, SyncJob
from app.services.calculation import prepare_calculation_batch, resolve_source_key
from app.services.event_coverage import evaluate_event_coverage
from app.services.facility_data import facility_data_status, queue_facility_load
from app.services.geography import descendants
from app.services.geometry import snapshot_map_features
from app.services.period_rollup import ensure_period_rollups
from app.services.publishing_documents import _trend_drawing
from tests.conftest import auth_header, login, query_dashboard
from tests.helpers import map_source, put_population, put_raw
from tests.test_period_rollup import FY_MONTHS, _current, _month_row, _programme, _unit


def test_invalid_month_never_becomes_a_short_valid_year(session):
    programme, unit = _programme(session), _unit(session, "PADER")
    map_source(session, programme.id, "ANC1", "TEST_ONLY_ANC1")
    rows = [_month_row(session, unit, programme, month, 10) for month in FY_MONTHS]
    rows[4].value_invalid = True
    session.flush()
    ensure_period_rollups(session, ["FY2025/26"])
    [year] = _current(session, unit, "FY2025/26")
    assert year.value is None and year.value_invalid
    assert len(year.provenance["components"]) == 12
    assert str(rows[4].id) in year.provenance["component_row_ids"]


def test_withdrawn_month_and_then_all_months_withdraw_derived_year(session):
    programme, unit = _programme(session), _unit(session, "PADER")
    map_source(session, programme.id, "ANC1", "TEST_ONLY_ANC1")
    rows = [_month_row(session, unit, programme, month, 10) for month in FY_MONTHS]
    ensure_period_rollups(session, ["FY2025/26"])
    assert len(_current(session, unit, "FY2025/26")) == 1
    rows[0].is_current = False
    session.flush()
    ensure_period_rollups(session, ["FY2025/26"])
    assert _current(session, unit, "FY2025/26") == []
    rows[0].is_current = True
    session.flush()
    ensure_period_rollups(session, ["FY2025/26"])
    for row in rows:
        row.is_current = False
    session.flush()
    ensure_period_rollups(session, ["FY2025/26"])
    assert _current(session, unit, "FY2025/26") == []


def test_other_district_rows_do_not_certify_an_unrequested_unit_key(session):
    programme = _programme(session)
    map_source(session, programme.id, "ANC1", "TEST_ONLY_ANC1")
    pader, kitgum = _unit(session, "PADER"), _unit(session, "KITGUM")
    for month in FY_MONTHS:
        _month_row(session, kitgum, programme, month, 10)
    _month_row(session, pader, programme, FY_MONTHS[0], 10)
    ensure_period_rollups(session, ["FY2025/26"])
    assert _current(session, pader, "FY2025/26") == []


def test_disabled_native_percentage_is_never_summed(session):
    programme, unit = _programme(session), _unit(session, "PADER")
    mapping = map_source(session, programme.id, "KMC_PERCENT", "TEST_ONLY_NATIVE")
    mapping.item_kind, mapping.enabled = "indicator", False
    for month in FY_MONTHS:
        _month_row(session, unit, programme, month, 50, key="KMC_PERCENT")
    session.flush()
    ensure_period_rollups(session, ["FY2025/26"])
    assert _current(session, unit, "FY2025/26", "KMC_PERCENT") == []


def test_no_district_rows_cannot_fall_back_to_facilities(session):
    """A region reads its district/city rows only (D-074): a facility row is never a substitute,
    and it does not even change what the region reports, so loading facilities on demand cannot
    alter a regional or national figure."""
    acholi = _unit(session, "ACHOLI")
    before = resolve_source_key(session, acholi, "FY2024/25", "ANC1")
    put_raw(session, _unit(session, "PADER_HC_III"), "FY2024/25", "ANC1", 100)
    resolved = resolve_source_key(session, acholi, "FY2024/25", "ANC1")
    assert resolved.value is None and resolved.status == "missing"
    assert (resolved.value, resolved.status) == (before.value, before.status)


def test_routine_calculation_batch_never_selects_event_rows(session, monkeypatch):
    from sqlalchemy import event

    statements = []

    def record(_connection, _cursor, statement, *_args):
        statements.append(statement)

    event.listen(session.get_bind(), "before_cursor_execute", record)
    try:
        batch = prepare_calculation_batch(
            session, org_unit=_unit(session, "UG"), periods=["202506"], programme_codes=["MNCH"]
        )
        assert batch["events"] == []
        assert not any("FROM raw_event_snapshots" in sql for sql in statements)
    finally:
        event.remove(session.get_bind(), "before_cursor_execute", record)


def test_closed_facility_remains_in_its_dated_history(session):
    facility = _unit(session, "PADER_HC_III")
    facility.active, facility.valid_to = False, date(2025, 7, 31)
    session.flush()
    district = _unit(session, "PADER")
    assert facility in descendants(session, district, period="202506")
    assert facility not in descendants(session, district, period="202508")
    assert facility not in descendants(session, district)


def test_multi_form_mpdsr_zero_requires_all_requested_programs(session):
    programme = _programme(session, "MPDSR")
    mappings = []
    for index in range(4):
        mapping = EventFieldMapping(
            programme_id=programme.id,
            program_uid=f"TEST_PROGRAM_{index}",
            internal_semantic_field="death_date",
            event_type=f"TEST_FORM_{index}",
            mapping_version="TEST_MULTI",
            enabled=True,
        )
        session.add(mapping)
        mappings.append(mapping)
    session.flush()
    job = SyncJob(
        job_type="tracker",
        status="succeeded",
        programme_id=programme.id,
        org_unit_id=_unit(session, "UG").id,
        mapping_version="TEST_MULTI",
        finished_at=datetime.now(UTC),
        window_start=date(2024, 7, 1),
        window_end=date(2025, 6, 30),
    )
    session.add(job)
    session.flush()
    kwargs = dict(org_unit=_unit(session, "PADER"), period="FY2024/25", required_end=date(2025, 6, 30))
    assert evaluate_event_coverage(session, **kwargs)["status"] == "unverified"
    job.coverage_manifest = {
        "complete": True,
        "program_uids": [m.program_uid for m in mappings],
        "mapping_ids": [str(m.id) for m in mappings],
    }
    session.flush()
    assert evaluate_event_coverage(session, **kwargs)["status"] == "verified"
    job.coverage_manifest = {**job.coverage_manifest, "missing_death_dates": 1}
    session.flush()
    assert evaluate_event_coverage(session, **kwargs)["coverage_failure"] == "missing_dates"
    job.coverage_manifest = {**job.coverage_manifest, "missing_death_dates": 0}
    job.coverage_manifest = {**job.coverage_manifest, "program_uids": [mappings[0].program_uid]}
    session.flush()
    assert evaluate_event_coverage(session, **kwargs)["status"] == "unverified"


def test_success_date_alone_does_not_mean_facilities_ready(session):
    programme = _programme(session)
    map_source(session, programme.id, "ANC1", "TEST_ONLY_ANC1")
    session.add(
        SyncJob(
            job_type="aggregate_facilities",
            status="succeeded",
            programme_id=programme.id,
            org_unit_id=_unit(session, "PADER").id,
            period_from="202506",
            period_to="202506",
            finished_at=datetime.now(UTC),
            mapping_version="v1",
        )
    )
    session.flush()
    status = facility_data_status(session, org_unit=_unit(session, "PADER"), period="202506", module="anc")
    assert status.start_job and status.months_loaded == 0


def test_facility_load_does_not_expand_to_its_district(session, monkeypatch):
    from app.models import User
    from app.services import facility_data

    programme = _programme(session)
    map_source(session, programme.id, "ANC1", "TEST_ONLY_ANC1")
    monkeypatch.setattr(facility_data, "ensure_extraction_enabled", lambda _: None)
    user = session.scalar(select(User).where(User.username == "paderhc3.user"))
    assert user is not None
    facility = _unit(session, "PADER_HC_III")
    status = queue_facility_load(session, user=user, org_unit=facility, period="202506", module="anc")
    assert session.get(SyncJob, UUID(status.job_id)).org_unit_id == facility.id


def test_prepared_repeat_is_read_only_and_skips_calculation(client, session, monkeypatch):
    from app.services import dashboard

    unit = _unit(session, "ACHOLI")
    put_population(session, unit, 2024, 1_000_000)
    put_raw(session, unit, "FY2024/25", "ANC1", 48_000)
    headers = auth_header(login(client, "admin.user"))
    first = query_dashboard(client, headers, unit.id, request_key="october-first").json()
    before = session.scalar(select(func.count(CalculationRun.id)))

    def fail(*_args, **_kwargs):
        raise AssertionError("a prepared view must not calculate")

    monkeypatch.setattr(dashboard, "evaluate_module", fail)
    second = query_dashboard(client, headers, unit.id, request_key="october-second")
    assert second.status_code == 200
    assert second.json()["analysis_snapshot_id"] == first["analysis_snapshot_id"]
    prepared = client.get(
        "/analytics/dashboard/prepared", params={"org_unit_id": str(unit.id), "period": "FY2024/25", "module": "anc"}
    )
    assert prepared.json()["ready"] is True
    assert session.scalar(select(func.count(CalculationRun.id))) == before


def test_bulk_source_edit_invalidates_prepared_product(client, session):
    unit = _unit(session, "ACHOLI")
    put_raw(session, unit, "FY2024/25", "ANC1", 48_000)
    headers = auth_header(login(client, "admin.user"))
    first = query_dashboard(client, headers, unit.id, request_key="october-before").json()
    revision = session.scalar(select(AnalyticsRevision.revision))
    from app.models import RawAggregateValue

    session.execute(update(RawAggregateValue).where(RawAggregateValue.org_unit_id == unit.id).values(value=49_000))
    session.commit()
    assert session.scalar(select(AnalyticsRevision.revision)) > revision
    prepared = client.get("/analytics/dashboard/prepared", params={"org_unit_id": str(unit.id), "period": "FY2024/25"})
    assert prepared.json()["ready"] is False
    second = query_dashboard(client, headers, unit.id, request_key="october-after").json()
    assert second["analysis_snapshot_id"] != first["analysis_snapshot_id"]


def test_pdf_trend_has_separate_segments_at_missing_months():
    from types import SimpleNamespace

    from reportlab.graphics.charts.lineplots import LinePlot

    view = SimpleNamespace(
        trend_indicator={"indicator_code": "X", "unit": "%"},
        trends=[
            {"period": month, "values": {"X": {"raw_value": value}}}
            for month, value in [("202501", 10), ("202502", None), ("202503", 30)]
        ],
    )
    drawing = _trend_drawing(view)
    plot = next(item for item in drawing.contents if isinstance(item, LinePlot))
    assert plot.data == [[(0, 10)], [(2, 30)]]


def test_snapshot_map_uses_pinned_boundary_not_new_replacement(session):
    from app.models import User

    unit = _unit(session, "PADER")
    old = Geometry(org_unit_id=unit.id, geojson={"type": "Point", "coordinates": [32, 2]}, geometry_kind="point")
    session.add(old)
    session.flush()
    old.valid_to = date(2025, 6, 29)
    session.flush()
    session.add(Geometry(org_unit_id=unit.id, geojson={"type": "Point", "coordinates": [33, 3]}, geometry_kind="point"))
    session.flush()
    block = {
        "selected_indicator": "X",
        "geometry_effective_date": "2025-06-30",
        "map_feature_org_unit_ids": [str(unit.id)],
        "geometry_versions": {str(unit.id): {"geometry_id": str(old.id)}},
    }
    user = session.scalar(select(User).where(User.username == "admin.user"))
    features = snapshot_map_features(
        session, user, map_block=block, value_rows=[{"org_unit_id": str(unit.id), "values": {"X": {"raw_value": 10}}}]
    )
    assert features["features"][0]["geometry"]["coordinates"] == [32, 2]


def test_pdf_names_are_literal_text_not_reportlab_markup():
    from types import SimpleNamespace

    from app.services.publishing_documents import _kpi_table, _pdf_styles

    name = '<img src="no-such-file.png"/> & literal text'
    table = _kpi_table(
        SimpleNamespace(kpis=[{"name": name, "raw_value": 10, "unit": "%", "status": "green"}]), _pdf_styles()
    )
    paragraph = table._cellvalues[0][0][0]
    assert paragraph.getPlainText() == name


def test_excel_source_text_cannot_become_an_executable_formula(client, session, tmp_path, monkeypatch):
    from app.models import Indicator
    from tests.test_publishing_visuals import _export, _national, _path

    indicator = session.scalar(select(Indicator).where(Indicator.code == "ANC1_COVERAGE"))
    indicator.name = "=1+1"
    session.commit()
    headers, unit, dashboard = _national(client, session, tmp_path, monkeypatch)
    book = load_workbook(_path(session, _export(client, headers, unit, dashboard, "excel")))
    cells = [cell for sheet in book for row in sheet for cell in row]
    assert any(cell.value == "=1+1" and cell.data_type == "s" for cell in cells)
    assert all(cell.data_type != "f" for cell in cells)


def test_preparation_cannot_claim_a_revision_that_changed_while_building(session):
    from app.models import User
    from app.services.prepared_dashboard import preparation_metadata

    programme, unit = _programme(session), _unit(session, "PADER")
    put_raw(session, unit, "202506", "ANC1", 10)
    prepare_calculation_batch(session, org_unit=unit, periods=["202506"], programme_codes=[programme.code])
    captured = session.info["product_input_revision"]
    put_raw(session, unit, "202506", "ANC4", 11)
    metadata = preparation_metadata(session, session.scalar(select(User).where(User.username == "admin.user")), "anc")
    assert metadata["input_revision"] == captured
    assert metadata["revision_changed_during_preparation"] is True
    assert metadata["signature"] is None


def test_invalid_peer_cannot_be_nonreporting_even_with_extraction_proof(session):
    from app.services.calculation import _child_cohort

    programme = _programme(session)
    parent = _unit(session, "ACHOLI")
    first, invalid_unit = _unit(session, "PADER"), _unit(session, "KITGUM")
    put_raw(session, first, "202506", "ANC1", 10)
    invalid = put_raw(session, invalid_unit, "202506", "ANC1", 11)
    invalid.value_invalid = True
    session.add(
        SyncJob(
            job_type="aggregate",
            status="succeeded",
            programme_id=programme.id,
            org_unit_id=parent.id,
            mapping_version="v1",
            period_from="202506",
            period_to="202506",
            coverage_manifest={
                "complete": True,
                "periods": ["202506"],
                "source_keys": ["ANC1"],
                "org_unit_ids": [str(first.id), str(invalid_unit.id)],
            },
        )
    )
    session.flush()
    result = _child_cohort(
        session, descendants(session, parent, period="202506"), "202506", "ANC1", programme.id, allow_partial=True
    )
    assert result.status == "incomplete" and "invalid peer" in result.reason.lower()


def test_pdf_maps_preserve_holes_and_render_facility_points():
    from types import SimpleNamespace

    from reportlab.graphics.shapes import Circle
    from reportlab.graphics.shapes import Path as DrawingPath

    from app.services.publishing_documents import _map_drawing

    shell = [[0, 0], [4, 0], [4, 4], [0, 4], [0, 0]]
    hole = [[1, 1], [3, 1], [3, 3], [1, 3], [1, 1]]
    view = SimpleNamespace(
        map_features=[
            {"geometry": {"type": "Polygon", "coordinates": [shell, hole]}},
            {"geometry": {"type": "Point", "coordinates": [2, 2]}},
        ]
    )
    drawing = _map_drawing(view, 100, 100)
    path = next(item for item in drawing.contents if isinstance(item, DrawingPath))
    assert len(path.points) == 20 and path.fillMode == 0
    assert sum(isinstance(item, Circle) for item in drawing.contents) == 1


def test_sensitive_mpdsr_snapshot_rechecks_event_permission(session):
    import pytest

    from app.domain.enums import ActionPermission
    from app.models import AnalysisSnapshot, User, UserProgrammeScope
    from app.services.analysis import load_snapshot
    from app.services.authorization import AuthorizationError, has_action

    user = session.scalar(select(User).where(User.username == "pader.focal"))
    assert not has_action(session, user, ActionPermission.VIEW_MPDSR_EVENTS)
    session.add(UserProgrammeScope(user_id=user.id, programme_id=_programme(session, "MPDSR").id, is_active=True))
    session.flush()
    snapshot = AnalysisSnapshot(
        user_id=user.id,
        org_unit_id=_unit(session, "PADER").id,
        period="FY2024/25",
        module="mpdsr",
        status="succeeded",
        view_hash="TEST_ONLY_SENSITIVE_SNAPSHOT",
        payload_json={"module_result": {"mpdsr": {"active_events": {"count": 2}}}},
    )
    session.add(snapshot)
    session.flush()
    with pytest.raises(AuthorizationError, match="Missing action permission"):
        load_snapshot(session, user=user, snapshot_id=snapshot.id)


def test_disabled_programme_cannot_survive_a_users_active_grant(session):
    from app.models import User
    from app.services.authorization import programme_scope_codes

    user = session.scalar(select(User).where(User.username == "pader.focal"))
    programme = _programme(session)
    assert "MNCH" in programme_scope_codes(session, user)
    programme.active = False
    session.flush()
    assert "MNCH" not in programme_scope_codes(session, user)
