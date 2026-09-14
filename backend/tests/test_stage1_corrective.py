from datetime import date
from uuid import UUID, uuid4

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import event, select
from sqlalchemy.orm import sessionmaker

from app.config import get_settings
from app.db.base import Base
from app.db.session import create_db_engine, reset_engine
from app.domain.enums import MappingSourceSystem, OrgUnitLevel
from app.main import app
from app.models import (
    AnalysisSnapshot,
    CalculationRun,
    Indicator,
    IndicatorVersion,
    OrgUnit,
    OrgUnitMapping,
    Programme,
    User,
)
from app.services.authorization import AuthorizationError
from app.services.calculation import select_indicator_versions_for_period
from app.services.geography import descendants, resolve_org_unit_by_dhis2_uid
from app.services.modules import evaluate_module
from app.services.seed import seed_reference_data
from tests.conftest import SEED_PASSWORD, auth_header, login, query_dashboard
from tests.helpers import put_event, put_population, put_raw


def _unit(session, code):
    return session.scalar(select(OrgUnit).where(OrgUnit.code == code))


def test_dashboard_snapshot_exists_in_a_new_session(tmp_path, monkeypatch):
    db_path = tmp_path / "persist.db"
    url = f"sqlite+pysqlite:///{db_path}"
    monkeypatch.setenv("DATABASE_URL", url)
    get_settings.cache_clear()
    reset_engine()
    engine = create_db_engine()
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)
    seed_session = factory()
    seed_reference_data(seed_session, SEED_PASSWORD)
    pader = seed_session.scalar(select(OrgUnit).where(OrgUnit.code == "PADER"))
    pader_id = pader.id
    put_population(seed_session, pader, 2024, 1_000_000, code="PERSIST_POP")
    put_raw(seed_session, pader, "FY2024/25", "ANC1", 40_000)
    seed_session.commit()
    seed_session.close()

    app.dependency_overrides.clear()
    with TestClient(app) as client:
        csrf = login(client, "pader.focal")
        dash = query_dashboard(client, auth_header(csrf), pader_id)
        assert dash.status_code == 201, dash.text
        snapshot_id = dash.json()["analysis_snapshot_id"]
        run_id = dash.json()["module_result"]["current_run_id"]

    other = factory()
    stored = other.get(AnalysisSnapshot, UUID(snapshot_id))
    assert stored is not None
    assert stored.payload_json["module_result"]["current_run_id"] == run_id
    assert stored.payload_json["module_result"]["indicators"]
    other.close()
    get_settings.cache_clear()
    reset_engine()


def test_failed_dashboard_leaves_no_completed_snapshot(tmp_path, monkeypatch):
    db_path = tmp_path / "fail.db"
    url = f"sqlite+pysqlite:///{db_path}"
    monkeypatch.setenv("DATABASE_URL", url)
    get_settings.cache_clear()
    reset_engine()
    engine = create_db_engine()
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)
    seed_session = factory()
    seed_reference_data(seed_session, SEED_PASSWORD)
    pader = seed_session.scalar(select(OrgUnit).where(OrgUnit.code == "PADER"))
    pader_id = pader.id
    seed_session.commit()
    seed_session.close()

    def boom(*_args, **_kwargs):
        raise RuntimeError("forced snapshot failure")

    monkeypatch.setattr("app.services.dashboard.persist_snapshot", boom)
    app.dependency_overrides.clear()
    with TestClient(app, raise_server_exceptions=False) as client:
        csrf = login(client, "pader.focal")
        response = query_dashboard(client, auth_header(csrf), pader_id)
        assert response.status_code >= 500

    other = factory()
    completed = other.scalars(select(AnalysisSnapshot)).all()
    assert completed == []
    assert other.scalars(select(CalculationRun)).all() == []
    other.close()
    get_settings.cache_clear()
    reset_engine()


def test_like_wildcard_codes_do_not_expand_subtrees(session):
    parent = OrgUnit(
        id=uuid4(),
        code="UNIT_ONE",
        name="Wildcard parent",
        level_type=OrgUnitLevel.DISTRICT.value,
        path="/UNIT_ONE",
        active=True,
    )
    child = OrgUnit(
        id=uuid4(),
        code="REAL_CHILD",
        name="Real child",
        level_type=OrgUnitLevel.FACILITY.value,
        parent_id=parent.id,
        path="/UNIT_ONE/REAL_CHILD",
        active=True,
    )
    decoy = OrgUnit(
        id=uuid4(),
        code="DECOY",
        name="Decoy from underscore wildcard",
        level_type=OrgUnitLevel.FACILITY.value,
        path="/UNITXONE/DECOY",
        active=True,
    )
    percent_parent = OrgUnit(
        id=uuid4(),
        code="PCT%ROOT",
        name="Percent parent",
        level_type=OrgUnitLevel.DISTRICT.value,
        path="/PCT%ROOT",
        active=True,
    )
    session.add_all([parent, child, decoy, percent_parent])
    session.flush()
    codes = {unit.code for unit in descendants(session, parent, include_self=False)}
    assert codes == {"REAL_CHILD"}
    assert "DECOY" not in codes


def test_org_unit_mapping_is_effective_date_aware(session):
    pader = _unit(session, "PADER")
    kitgum = _unit(session, "KITGUM")
    uid = "TEST_UID_OU_REMAP"
    session.add(
        OrgUnitMapping(
            org_unit_id=pader.id,
            source_system=MappingSourceSystem.DHIS2.value,
            external_uid=uid,
            valid_from=date(2020, 1, 1),
            valid_to=date(2023, 12, 31),
        )
    )
    session.add(
        OrgUnitMapping(
            org_unit_id=kitgum.id,
            source_system=MappingSourceSystem.DHIS2.value,
            external_uid=uid,
            valid_from=date(2024, 1, 1),
            valid_to=None,
        )
    )
    session.flush()
    assert resolve_org_unit_by_dhis2_uid(session, uid, as_of=date(2022, 6, 30)).id == pader.id
    assert resolve_org_unit_by_dhis2_uid(session, uid, as_of=date(2025, 1, 1)).id == kitgum.id


def test_overlapping_org_unit_mappings_are_rejected(session):
    pader = _unit(session, "PADER")
    kitgum = _unit(session, "KITGUM")
    uid = "TEST_UID_OU_OVERLAP"
    session.add(
        OrgUnitMapping(
            org_unit_id=pader.id,
            source_system=MappingSourceSystem.DHIS2.value,
            external_uid=uid,
            valid_from=date(2024, 1, 1),
            valid_to=date(2025, 12, 31),
        )
    )
    session.add(
        OrgUnitMapping(
            org_unit_id=kitgum.id,
            source_system=MappingSourceSystem.DHIS2.value,
            external_uid=uid,
            valid_from=date(2025, 1, 1),
            valid_to=None,
        )
    )
    session.flush()
    with pytest.raises(AuthorizationError):
        resolve_org_unit_by_dhis2_uid(session, uid, as_of=date(2025, 6, 1))


def test_mpdsr_ignores_foreign_program_events(session):
    acholi = _unit(session, "ACHOLI")
    admin = session.scalar(select(User).where(User.username == "admin.user"))
    epi = session.scalar(select(Programme).where(Programme.code == "EPI"))
    put_event(
        session,
        acholi,
        event_uid="TEST_UID_MPDSR_OK",
        data_values={"event_type": "maternal_notification"},
        death_date=date(2024, 8, 1),
        notification_date=date(2024, 8, 2),
        status="COMPLETED",
    )
    put_event(
        session,
        acholi,
        event_uid="TEST_UID_FOREIGN_PROGRAM",
        programme_id=epi.id,
        program_uid="TEST_UID_OTHER_TRACKER",
        data_values={"event_type": "maternal_notification"},
        death_date=date(2024, 8, 1),
        notification_date=date(2024, 8, 2),
        status="COMPLETED",
    )
    result = evaluate_module(
        session,
        user=admin,
        org_unit_id=acholi.id,
        period="FY2024/25",
        module="mpdsr",
        include_children=False,
    )
    by_code = {row["indicator_code"]: row for row in result["indicators"]}
    assert by_code["MATERNAL_NOTIFIED_COUNT"]["raw_value"] == 1


def test_historical_formula_version_is_selected(session):
    indicator = session.scalar(select(Indicator).where(Indicator.code == "ANC1_COVERAGE"))
    current = session.scalar(
        select(IndicatorVersion).where(
            IndicatorVersion.indicator_id == indicator.id,
            IndicatorVersion.is_current.is_(True),
        )
    )
    current.valid_from = date(2025, 7, 1)
    current.valid_to = None
    historic = IndicatorVersion(
        indicator_id=indicator.id,
        formula_version="historic-test",
        numerator_definition=current.numerator_definition,
        denominator_type=current.denominator_type,
        denominator_coefficient=current.denominator_coefficient,
        multiplier=current.multiplier,
        unit=current.unit,
        display_precision=current.display_precision,
        direction=current.direction,
        formula_spec=current.formula_spec,
        classification_spec=current.classification_spec,
        valid_from=date(2020, 7, 1),
        valid_to=date(2025, 6, 30),
        is_current=False,
    )
    session.add(historic)
    session.flush()
    selected = select_indicator_versions_for_period(session, "FY2024/25")
    chosen = next(row for row in selected if row.indicator_id == indicator.id)
    assert chosen.formula_version == "historic-test"
    later = select_indicator_versions_for_period(session, "FY2025/26")
    current_chosen = next(row for row in later if row.indicator_id == indicator.id)
    assert current_chosen.id == current.id


def test_export_and_ai_require_matching_snapshot(client, session):
    pader = _unit(session, "PADER")
    put_population(session, pader, 2024, 1_000_000, code="SNAP_POP")
    put_raw(session, pader, "FY2024/25", "ANC1", 40_000)
    session.commit()
    csrf = login(client, "pader.focal")
    headers = auth_header(csrf)
    missing = client.post(
        "/exports/excel",
        json={"org_unit_id": str(pader.id), "period": "FY2024/25", "module": "anc"},
        headers=headers,
    )
    assert missing.status_code == 422
    dash = query_dashboard(client, headers, pader.id).json()
    mismatch = client.post(
        "/ai/findings",
        json={
            "org_unit_id": str(pader.id),
            "period": "FY2023/24",
            "module": "anc",
            "analysis_snapshot_id": dash["analysis_snapshot_id"],
            "view_hash": dash["view_hash"],
        },
        headers=headers,
    )
    assert mismatch.status_code == 422
    conflict = client.post(
        "/exports/excel",
        json={
            "org_unit_id": str(pader.id),
            "period": "FY2024/25",
            "module": "anc",
            "analysis_snapshot_id": dash["analysis_snapshot_id"],
            "view_hash": "0" * 64,
        },
        headers=headers,
    )
    assert conflict.status_code == 409


def test_historical_export_does_not_follow_later_raw_changes(client, session, tmp_path, monkeypatch):
    from openpyxl import load_workbook

    from app.config import get_settings

    pader = _unit(session, "PADER")
    put_population(session, pader, 2024, 1_000_000, code="HIST_POP")
    put_raw(session, pader, "FY2024/25", "ANC1", 47_700)
    session.commit()
    monkeypatch.setattr(get_settings(), "export_dir", str(tmp_path))
    csrf = login(client, "pader.focal")
    headers = auth_header(csrf)
    dash = query_dashboard(client, headers, pader.id).json()
    original = next(
        row["raw_value"]
        for row in dash["module_result"]["indicators"]
        if row["indicator_code"] == "ANC1_COVERAGE"
    )
    put_raw(session, pader, "FY2024/25", "ANC1", 10_000)
    session.commit()
    created = client.post(
        "/exports/excel",
        json={
            "org_unit_id": str(pader.id),
            "period": "FY2024/25",
            "module": "anc",
            "analysis_snapshot_id": dash["analysis_snapshot_id"],
            "view_hash": dash["view_hash"],
        },
        headers=headers,
    )
    assert created.status_code == 202, created.text
    path = tmp_path / f"{created.json()['job_id']}.xlsx"
    book = load_workbook(path)
    names = {row[0]: row for row in book["Scorecard"].iter_rows(min_row=2, values_only=True)}
    anc1_name = next(
        row["name"] for row in dash["module_result"]["indicators"] if row["indicator_code"] == "ANC1_COVERAGE"
    )
    assert names[anc1_name][1] == original


def test_module_batch_keeps_query_count_bounded(session):
    pader = _unit(session, "PADER")
    admin = session.scalar(select(User).where(User.username == "admin.user"))
    put_population(session, pader, 2024, 1_000_000, code="QCOUNT_POP")
    put_raw(session, pader, "FY2024/25", "ANC1", 10_000)
    engine = session.get_bind()
    count = {"n": 0}

    def _on_execute(*_args, **_kwargs):
        count["n"] += 1

    event.listen(engine, "after_cursor_execute", _on_execute)
    try:
        evaluate_module(
            session,
            user=admin,
            org_unit_id=pader.id,
            period="FY2024/25",
            module="anc",
            include_children=True,
        )
    finally:
        event.remove(engine, "after_cursor_execute", _on_execute)
    assert count["n"] < get_settings().query_count_warn


def test_concurrent_dashboards_do_not_cross_contaminate(tmp_path, monkeypatch):
    from threading import Thread

    db_path = tmp_path / "concurrent.db"
    url = f"sqlite+pysqlite:///{db_path}"
    monkeypatch.setenv("DATABASE_URL", url)
    get_settings.cache_clear()
    reset_engine()
    engine = create_db_engine()
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)
    seed_session = factory()
    seed_reference_data(seed_session, SEED_PASSWORD)
    pader = seed_session.scalar(select(OrgUnit).where(OrgUnit.code == "PADER"))
    acholi = seed_session.scalar(select(OrgUnit).where(OrgUnit.code == "ACHOLI"))
    pader_id, acholi_id = pader.id, acholi.id
    put_population(seed_session, pader, 2024, 1_000_000, code="CONC_PADER")
    put_population(seed_session, acholi, 2024, 2_000_000, code="CONC_ACHOLI")
    put_raw(seed_session, pader, "FY2024/25", "ANC1", 40_000)
    put_raw(seed_session, acholi, "FY2024/25", "ANC1", 80_000)
    seed_session.commit()
    seed_session.close()
    app.dependency_overrides.clear()
    results: dict[str, dict] = {}

    def _load(username: str, org_id, key: str) -> None:
        with TestClient(app) as client:
            csrf = login(client, username)
            response = query_dashboard(client, auth_header(csrf), org_id)
            assert response.status_code == 201, response.text
            results[key] = response.json()

    workers = [
        Thread(target=_load, args=("pader.focal", pader_id, "pader")),
        Thread(target=_load, args=("acholi.analyst", acholi_id, "acholi")),
    ]
    for worker in workers:
        worker.start()
    for worker in workers:
        worker.join(timeout=60)
    assert set(results) == {"pader", "acholi"}
    assert results["pader"]["analysis_snapshot_id"] != results["acholi"]["analysis_snapshot_id"]
    other = factory()
    first = other.get(AnalysisSnapshot, UUID(results["pader"]["analysis_snapshot_id"]))
    second = other.get(AnalysisSnapshot, UUID(results["acholi"]["analysis_snapshot_id"]))
    assert first is not None and second is not None
    assert first.org_unit_id == pader_id
    assert second.org_unit_id == acholi_id
    other.close()
    get_settings.cache_clear()
    reset_engine()


def test_export_stays_queued_until_worker_runs(client, session, tmp_path, monkeypatch):
    from app.domain.enums import JobStatus
    from app.models import ExportJob
    from app.services.publishing import generate_export_artifact

    pader = _unit(session, "PADER")
    put_population(session, pader, 2024, 1_000_000, code="QUEUE_POP")
    put_raw(session, pader, "FY2024/25", "ANC1", 40_000)
    session.commit()
    monkeypatch.setattr(get_settings(), "export_dir", str(tmp_path))
    monkeypatch.setattr(get_settings(), "export_eager", False)
    csrf = login(client, "pader.focal")
    headers = auth_header(csrf)
    dash = query_dashboard(client, headers, pader.id).json()
    created = client.post(
        "/exports/excel",
        json={
            "org_unit_id": str(pader.id),
            "period": "FY2024/25",
            "module": "anc",
            "analysis_snapshot_id": dash["analysis_snapshot_id"],
            "view_hash": dash["view_hash"],
        },
        headers=headers,
    )
    assert created.status_code == 202
    job_id = created.json()["job_id"]
    assert created.json()["status"] == JobStatus.QUEUED.value
    meta = client.get(f"/exports/jobs/{job_id}", headers=headers)
    assert meta.json()["downloadable"] is False
    generate_export_artifact(session, UUID(job_id))
    session.commit()
    ready = client.get(f"/exports/jobs/{job_id}", headers=headers)
    assert ready.json()["downloadable"] is True
    stored = session.get(ExportJob, UUID(job_id))
    assert stored is not None and stored.checksum
