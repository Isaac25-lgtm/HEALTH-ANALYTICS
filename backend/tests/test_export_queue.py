"""Export dispatch through Celery, idempotency, durable failure and bounded retries (SQLite file DB).

Celery publishing uses kombu's in-process ``memory://`` transport, so real task messages are
published and consumed without a network broker. PostgreSQL two-connection behaviour is covered
in ``test_postgres_export_queue.py``.
"""

from __future__ import annotations

import json
import time
from pathlib import Path
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from kombu.exceptions import OperationalError
from sqlalchemy import func, select, update
from sqlalchemy.orm import sessionmaker

from app.config import Settings, get_settings, validate_runtime_settings
from app.db.base import Base
from app.db.session import create_db_engine, reset_engine
from app.domain.enums import JobStatus
from app.main import app
from app.models import AuditLog, ExportJob, OrgUnit, RawAggregateValue
from app.services import export_jobs, publishing
from app.services.seed import seed_reference_data
from app.workers.celery_app import EXPORT_TASK_NAME, celery_app
from tests.conftest import SEED_PASSWORD, auth_header, download_export, login, query_dashboard
from tests.helpers import put_population, put_raw


def _purge_export_queue() -> None:
    with celery_app.connection_for_write() as connection:
        try:
            connection.default_channel.queue_purge(get_settings().export_queue_name)
        except Exception:  # noqa: BLE001 - the queue may not exist yet
            pass


def _drain_export_messages() -> list:
    messages = []
    with celery_app.connection_for_read() as connection:
        queue = connection.SimpleQueue(get_settings().export_queue_name, no_ack=True)
        try:
            while True:
                messages.append(queue.get(block=False))
        except queue.Empty:
            pass
        finally:
            queue.close()
    return messages


@pytest.fixture()
def queue_db(tmp_path, monkeypatch):
    """A file database reached through the real get_db dependency, in queue (non-eager) mode."""
    monkeypatch.setenv("DATABASE_URL", f"sqlite+pysqlite:///{tmp_path / 'queue.db'}")
    monkeypatch.setenv("EXPORT_DIR", str(tmp_path / "artifacts"))
    monkeypatch.setenv("EXPORT_EAGER", "false")
    get_settings.cache_clear()
    reset_engine()
    engine = create_db_engine()
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)
    with factory() as seed:
        seed_reference_data(seed, SEED_PASSWORD)
        pader = seed.scalar(select(OrgUnit).where(OrgUnit.code == "PADER"))
        put_population(seed, pader, 2024, 1_000_000, code="QUEUE_POP")
        put_raw(seed, pader, "FY2024/25", "ANC1", 40_000)
        seed.commit()
    app.dependency_overrides.clear()
    _purge_export_queue()
    try:
        yield engine, factory, tmp_path / "artifacts"
    finally:
        _purge_export_queue()
        engine.dispose()
        get_settings.cache_clear()
        reset_engine()


def _pader_id(factory) -> UUID:
    with factory() as session:
        return session.scalar(select(OrgUnit.id).where(OrgUnit.code == "PADER"))


def _export_body(pader_id, dash) -> dict:
    return {
        "org_unit_id": str(pader_id),
        "period": "FY2024/25",
        "module": "anc",
        "analysis_snapshot_id": dash["analysis_snapshot_id"],
        "view_hash": dash["view_hash"],
    }


def _job(factory, job_id) -> ExportJob:
    with factory() as session:
        job = session.get(ExportJob, UUID(str(job_id)))
        session.expunge(job)
        return job


def _submit(client, headers, pader_id, kind: str = "excel"):
    dash = query_dashboard(client, headers, pader_id)
    assert dash.status_code == 201, dash.text
    snapshot = dash.json()
    return snapshot, client.post(f"/exports/{kind}", json=_export_body(pader_id, snapshot), headers=headers)


def _artifacts(directory: Path) -> list[Path]:
    return sorted(directory.iterdir()) if directory.exists() else []


def test_job_and_audit_are_committed_before_dispatch_and_message_is_published(queue_db, monkeypatch):
    _engine, factory, _artifact_dir = queue_db
    pader_id = _pader_id(factory)
    observed: dict = {}
    original = export_jobs.send_export_task

    def spy(job_id):
        # An independent connection must already see the job and its audit record.
        with factory() as other:
            job = other.get(ExportJob, job_id)
            observed["status"] = job.status if job else None
            observed["audit"] = other.scalar(
                select(func.count())
                .select_from(AuditLog)
                .where(AuditLog.resource_id == str(job_id), AuditLog.action == "export_excel_requested")
            )
        return original(job_id)

    monkeypatch.setattr(export_jobs, "send_export_task", spy)
    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        _snapshot, created = _submit(client, headers, pader_id)
    assert created.status_code == 202, created.text
    body = created.json()
    assert body["status"] == JobStatus.QUEUED.value
    assert body["dispatch_state"] == export_jobs.DISPATCH_DISPATCHED
    assert observed == {"status": JobStatus.QUEUED.value, "audit": 1}
    job = _job(factory, body["job_id"])
    assert job.celery_task_id and job.dispatched_at is not None and job.file_path is None
    messages = _drain_export_messages()
    assert len(messages) == 1
    assert messages[0].headers["task"] == EXPORT_TASK_NAME
    assert messages[0].headers["id"] == job.celery_task_id
    assert messages[0].payload[0] == [body["job_id"]]


def test_queue_mode_never_generates_in_the_request(queue_db):
    _engine, factory, artifact_dir = queue_db
    pader_id = _pader_id(factory)
    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        _snapshot, created = _submit(client, headers, pader_id)
        meta = client.get(f"/exports/jobs/{created.json()['job_id']}", headers=headers).json()
    assert meta["status"] == JobStatus.QUEUED.value
    assert meta["downloadable"] is False
    assert _artifacts(artifact_dir) == []


def test_eager_generation_requires_explicit_development_or_test(monkeypatch):
    monkeypatch.delenv("EXPORT_EAGER", raising=False)
    assert Settings(_env_file=None).export_eager is False
    production = Settings(
        _env_file=None,
        app_env="production",
        export_eager=True,
        database_url="postgresql+psycopg://hpip:owner-supplied@db:5432/hpip",
    )
    assert any("EXPORT_EAGER" in error for error in validate_runtime_settings(production))
    for env in ("development", "test"):
        assert not any(
            "EXPORT_EAGER" in error
            for error in validate_runtime_settings(Settings(_env_file=None, app_env=env, export_eager=True))
        )


def test_eager_setting_outside_development_fails_closed_without_generating(queue_db, monkeypatch):
    _engine, factory, artifact_dir = queue_db
    pader_id = _pader_id(factory)
    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        _snapshot, created = _submit(client, headers, pader_id)
    job_id = UUID(created.json()["job_id"])
    with factory() as session:
        session.execute(update(ExportJob).where(ExportJob.id == job_id).values(status="queued"))
        session.commit()
        monkeypatch.setattr(get_settings(), "export_eager", True)
        monkeypatch.setattr(get_settings(), "app_env", "staging")
        job = export_jobs.start_export_job(session, job_id)
        assert job.status == JobStatus.FAILED.value
        assert job.error_code == export_jobs.QUEUE_UNCONFIGURED
    assert _artifacts(artifact_dir) == []


def test_broker_failure_is_honest_retryable_and_recovers_without_duplicates(queue_db, monkeypatch):
    _engine, factory, _artifact_dir = queue_db
    pader_id = _pader_id(factory)

    def broker_down(*_args, **_kwargs):
        raise OperationalError("connection refused to redis://:s3cr3t@broker.internal:6379/1")

    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        dash = query_dashboard(client, headers, pader_id).json()
        with monkeypatch.context() as patch:
            patch.setattr(celery_app, "send_task", broker_down)
            failed = client.post("/exports/excel", json=_export_body(pader_id, dash), headers=headers)
        assert failed.status_code == 503, failed.text
        detail = failed.json()["detail"]
        assert detail["code"] == export_jobs.ENQUEUE_FAILED
        assert detail["retryable"] is True
        encoded = json.dumps(failed.json())
        assert "s3cr3t" not in encoded and "redis://" not in encoded and "refused" not in encoded
        job = _job(factory, detail["job_id"])
        assert job.status == JobStatus.FAILED.value
        assert job.dispatch_state == export_jobs.DISPATCH_FAILED
        assert job.active_key is not None
        meta = client.get(f"/exports/jobs/{detail['job_id']}", headers=headers).json()
        assert meta["retryable"] is True and meta["permanent_failure"] is False
        retried = client.post("/exports/excel", json=_export_body(pader_id, dash), headers=headers)
    assert retried.status_code == 202, retried.text
    assert retried.json()["job_id"] == detail["job_id"]
    assert retried.json()["reused"] is True
    assert retried.json()["dispatch_state"] == export_jobs.DISPATCH_DISPATCHED
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(ExportJob)) == 1
    assert len(_drain_export_messages()) == 1


def test_retry_endpoint_redispatches_an_enqueue_failure(queue_db, monkeypatch):
    _engine, factory, _artifact_dir = queue_db
    pader_id = _pader_id(factory)
    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        dash = query_dashboard(client, headers, pader_id).json()
        with monkeypatch.context() as patch:
            patch.setattr(celery_app, "send_task", lambda *a, **k: (_ for _ in ()).throw(OperationalError("down")))
            failed = client.post("/exports/excel", json=_export_body(pader_id, dash), headers=headers)
        job_id = failed.json()["detail"]["job_id"]
        forged = client.post(f"/exports/jobs/{job_id}/retry")
        assert forged.status_code == 403
        other = auth_header(login(client, "view.only"))
        assert client.post(f"/exports/jobs/{job_id}/retry", headers=other).status_code in {403, 404}
        headers = auth_header(login(client, "pader.focal"))
        retried = client.post(f"/exports/jobs/{job_id}/retry", headers=headers)
        assert retried.status_code == 202, retried.text
        again = client.post(f"/exports/jobs/{job_id}/retry", headers=headers)
    assert retried.json()["dispatch_state"] == export_jobs.DISPATCH_DISPATCHED
    assert again.status_code == 409
    assert again.json()["detail"]["code"] == "export_retry_not_permitted"


def test_missing_broker_configuration_fails_closed(queue_db, monkeypatch):
    _engine, factory, _artifact_dir = queue_db
    pader_id = _pader_id(factory)
    calls: list = []
    monkeypatch.setattr(celery_app, "send_task", lambda *a, **k: calls.append(a))
    monkeypatch.setenv("CELERY_BROKER_URL", "")
    get_settings.cache_clear()
    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        _snapshot, created = _submit(client, headers, pader_id)
    assert created.status_code == 503
    assert created.json()["detail"]["code"] == export_jobs.QUEUE_UNCONFIGURED
    assert calls == []


def test_idempotent_resubmission_returns_one_job_one_message_and_one_artifact(queue_db):
    _engine, factory, artifact_dir = queue_db
    pader_id = _pader_id(factory)
    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        dash = query_dashboard(client, headers, pader_id).json()
        first = client.post("/exports/excel", json=_export_body(pader_id, dash), headers=headers)
        second = client.post("/exports/excel", json=_export_body(pader_id, dash), headers=headers)
        assert first.json()["job_id"] == second.json()["job_id"]
        assert second.json()["reused"] is True
        assert len(_drain_export_messages()) == 1
        from app.workers.tasks import run_generate_export_job

        assert run_generate_export_job(first.json()["job_id"]).outcome == "succeeded"
        third = client.post("/exports/excel", json=_export_body(pader_id, dash), headers=headers)
        other_type = client.post("/exports/report", json=_export_body(pader_id, dash), headers=headers)
    assert third.json()["job_id"] == first.json()["job_id"]
    assert third.json()["status"] == JobStatus.SUCCEEDED.value
    assert other_type.json()["job_id"] != first.json()["job_id"]
    assert [path.name for path in _artifacts(artifact_dir)] == [f"{first.json()['job_id']}.xlsx"]
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(ExportJob).where(ExportJob.export_type == "excel")) == 1


def test_concurrent_insert_race_resolves_to_the_existing_live_job(queue_db, monkeypatch):
    _engine, factory, _artifact_dir = queue_db
    pader_id = _pader_id(factory)
    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        dash = query_dashboard(client, headers, pader_id).json()
        first = client.post("/exports/excel", json=_export_body(pader_id, dash), headers=headers)
        real_lookup = export_jobs._live_job
        calls = {"n": 0}

        def stale_lookup(session, key):
            # The first lookup misses, as it would for a request that raced the first insert.
            calls["n"] += 1
            return None if calls["n"] == 1 else real_lookup(session, key)

        monkeypatch.setattr(export_jobs, "_live_job", stale_lookup)
        second = client.post("/exports/excel", json=_export_body(pader_id, dash), headers=headers)
    assert second.status_code == 202, second.text
    assert second.json()["job_id"] == first.json()["job_id"]
    with factory() as session:
        assert session.scalar(select(func.count()).select_from(ExportJob)) == 1


def _queued_job(factory, client, headers, pader_id) -> tuple[str, dict]:
    dash, created = _submit(client, headers, pader_id)
    assert created.status_code == 202, created.text
    _purge_export_queue()
    return created.json()["job_id"], dash


def test_worker_retry_moves_failed_job_to_succeeded(queue_db, monkeypatch):
    from app.workers.tasks import generate_export_job_task

    _engine, factory, artifact_dir = queue_db
    pader_id = _pader_id(factory)
    calls = {"n": 0}
    real_writer = publishing.ARTIFACT_WRITERS["excel"]

    def flaky(dashboard, package, destination):
        calls["n"] += 1
        if calls["n"] == 1:
            raise OSError(f"transient disk failure writing {destination}")
        real_writer(dashboard, package, destination)

    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        job_id, _dash = _queued_job(factory, client, headers, pader_id)
        monkeypatch.setitem(publishing.ARTIFACT_WRITERS, "excel", flaky)
        result = generate_export_job_task.apply(args=[job_id])
        meta = client.get(f"/exports/jobs/{job_id}", headers=headers).json()
    assert result.get()["outcome"] == "succeeded"
    job = _job(factory, job_id)
    assert job.status == JobStatus.SUCCEEDED.value
    assert job.attempt_count == 2
    assert job.error_code is None and job.retry_scheduled is False and job.checksum
    assert meta["downloadable"] is True
    with factory() as session:
        actions = session.scalars(
            select(AuditLog.action).where(AuditLog.resource_id == job_id).order_by(AuditLog.created_at)
        ).all()
    assert actions.count("export_failed") == 1 and actions.count("export_generated") == 1
    assert [path.name for path in _artifacts(artifact_dir)] == [f"{job_id}.xlsx"]


def test_permanent_failure_is_bounded_visible_and_safe(queue_db, monkeypatch, tmp_path):
    from app.workers.tasks import generate_export_job_task

    _engine, factory, artifact_dir = queue_db
    pader_id = _pader_id(factory)

    def broken(_dashboard, _package, destination):
        raise OSError(f"disk failure while writing {destination} secret-token-123")

    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        job_id, dash = _queued_job(factory, client, headers, pader_id)
        with monkeypatch.context() as patch:
            patch.setitem(publishing.ARTIFACT_WRITERS, "excel", broken)
            outcome = generate_export_job_task.apply(args=[job_id]).get()
        meta = client.get(f"/exports/jobs/{job_id}", headers=headers)
        listing = client.get("/exports/jobs", headers=headers).json()
        retry = client.post(f"/exports/jobs/{job_id}/retry", headers=headers)
        resubmitted = client.post("/exports/excel", json=_export_body(pader_id, dash), headers=headers)
    assert outcome["outcome"] == "failed"
    job = _job(factory, job_id)
    assert job.status == JobStatus.FAILED.value
    assert job.attempt_count == get_settings().export_max_attempts == 3
    assert job.error_code == export_jobs.GENERATION_FAILED
    assert job.retry_scheduled is False and job.active_key is None and job.finished_at is not None
    body = meta.json()
    assert body["permanent_failure"] is True and body["retryable"] is False
    assert body["error_message"] == export_jobs.SAFE_ERROR_MESSAGES[export_jobs.GENERATION_FAILED]
    encoded = json.dumps([body, listing])
    for leaked in ("disk failure", "secret-token-123", str(tmp_path), "Traceback", "OSError"):
        assert leaked not in encoded
    assert job_id in {row["job_id"] for row in listing["jobs"]}
    assert retry.status_code == 409
    assert resubmitted.status_code == 202
    assert resubmitted.json()["job_id"] != job_id
    assert _job(factory, job_id).status == JobStatus.FAILED.value
    assert not [path for path in _artifacts(artifact_dir) if path.name.endswith(".partial")]


def test_duplicate_delivery_does_not_regenerate(queue_db):
    from app.workers.tasks import run_generate_export_job

    _engine, factory, artifact_dir = queue_db
    pader_id = _pader_id(factory)
    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        job_id, _dash = _queued_job(factory, client, headers, pader_id)
    first = run_generate_export_job(job_id)
    checksum = _job(factory, job_id).checksum
    stat = (artifact_dir / f"{job_id}.xlsx").stat()
    second = run_generate_export_job(job_id)
    assert first.outcome == "succeeded"
    assert second.outcome == "skipped" and second.reason == "not_claimable"
    assert _job(factory, job_id).checksum == checksum
    assert (artifact_dir / f"{job_id}.xlsx").stat().st_mtime_ns == stat.st_mtime_ns
    assert [path.name for path in _artifacts(artifact_dir)] == [f"{job_id}.xlsx"]
    with factory() as session:
        generated = session.scalar(
            select(func.count())
            .select_from(AuditLog)
            .where(AuditLog.resource_id == job_id, AuditLog.action == "export_generated")
        )
    assert generated == 1


def test_stale_claim_is_taken_over_and_the_lost_claim_is_discarded(queue_db):
    _engine, factory, artifact_dir = queue_db
    pader_id = _pader_id(factory)
    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        job_id, _dash = _queued_job(factory, client, headers, pader_id)
    job_uuid = UUID(job_id)
    with factory() as session:
        assert export_jobs._claim(session, job_uuid, "stalled-worker-token")
        session.execute(
            update(ExportJob).where(ExportJob.id == job_uuid).values(claimed_at=export_jobs._now().replace(year=2000))
        )
        session.commit()
    takeover = export_jobs.process_export_job(job_uuid, factory=factory)
    assert takeover.outcome == "succeeded"
    holder: list[Path] = []
    with factory() as session:
        late = export_jobs._generate(session, job_uuid, "stalled-worker-token", holder)
    assert late.outcome == "skipped" and late.reason == "claim_lost"
    assert holder and not any(path.exists() for path in holder)
    job = _job(factory, job_id)
    assert job.status == JobStatus.SUCCEEDED.value and job.attempt_count == 2
    assert [path.name for path in _artifacts(artifact_dir)] == [f"{job_id}.xlsx"]


def test_failure_after_publication_step_rolls_back_work_in_progress(queue_db, monkeypatch):
    _engine, factory, artifact_dir = queue_db
    pader_id = _pader_id(factory)
    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        job_id, _dash = _queued_job(factory, client, headers, pader_id)

    def replace_fails(_src, _dst):
        raise PermissionError("cannot publish artifact")

    monkeypatch.setattr(export_jobs.os, "replace", replace_fails)
    attempt = export_jobs.process_export_job(UUID(job_id), factory=factory, auto_retry=False)
    assert attempt.outcome == "failed"
    job = _job(factory, job_id)
    # The success UPDATE and its audit row were in the rolled-back transaction.
    assert job.status == JobStatus.FAILED.value and job.file_path is None and job.checksum is None
    with factory() as session:
        actions = session.scalars(select(AuditLog.action).where(AuditLog.resource_id == job_id)).all()
    assert "export_generated" not in actions and "export_failed" in actions
    assert _artifacts(artifact_dir) == []


def test_worker_uses_the_committed_snapshot_not_later_raw_changes(queue_db):
    from openpyxl import load_workbook

    from app.workers.tasks import run_generate_export_job

    _engine, factory, artifact_dir = queue_db
    pader_id = _pader_id(factory)
    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        job_id, dash = _queued_job(factory, client, headers, pader_id)
    original = next(
        row["raw_value"] for row in dash["module_result"]["indicators"] if row["indicator_code"] == "ANC1_COVERAGE"
    )
    with factory() as session:
        pader = session.get(OrgUnit, pader_id)
        session.execute(update(RawAggregateValue).values(is_current=False))
        put_raw(session, pader, "FY2024/25", "ANC1", 10_000)
        session.commit()
    assert run_generate_export_job(job_id).outcome == "succeeded"
    book = load_workbook(artifact_dir / f"{job_id}.xlsx")
    name = next(row["name"] for row in dash["module_result"]["indicators"] if row["indicator_code"] == "ANC1_COVERAGE")
    rows = {row[0]: row for row in book["Scorecard"].iter_rows(min_row=2, values_only=True)}
    assert rows[name][1] == original


def test_revoked_access_fails_permanently_without_artifact(queue_db):
    from app.models import User, UserGeographyScope
    from app.workers.tasks import run_generate_export_job

    _engine, factory, artifact_dir = queue_db
    pader_id = _pader_id(factory)
    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        job_id, _dash = _queued_job(factory, client, headers, pader_id)
    with factory() as session:
        user = session.scalar(select(User).where(User.username == "pader.focal"))
        for scope in session.scalars(select(UserGeographyScope).where(UserGeographyScope.user_id == user.id)).all():
            session.delete(scope)
        session.commit()
    attempt = run_generate_export_job(job_id)
    job = _job(factory, job_id)
    assert attempt.outcome == "failed"
    assert job.error_code in {export_jobs.NOT_AUTHORISED, export_jobs.SNAPSHOT_UNAVAILABLE}
    assert job.active_key is None and job.attempt_count == 1
    assert _artifacts(artifact_dir) == []


def test_in_process_celery_worker_consumes_the_queue(queue_db):
    from celery.contrib.testing.worker import start_worker

    _engine, factory, _artifact_dir = queue_db
    pader_id = _pader_id(factory)
    with TestClient(app) as client:
        headers = auth_header(login(client, "pader.focal"))
        _snapshot, created = _submit(client, headers, pader_id)
        job_id = created.json()["job_id"]
        # A duplicate message for the same job must not create a second artifact.
        export_jobs.send_export_task(UUID(job_id))
        with start_worker(
            celery_app,
            pool="solo",
            perform_ping_check=False,
            queues=[get_settings().export_queue_name],
            shutdown_timeout=30,
        ):
            deadline = time.monotonic() + 60
            status = None
            while time.monotonic() < deadline:
                status = _job(factory, job_id).status
                if status in {JobStatus.SUCCEEDED.value, JobStatus.FAILED.value}:
                    break
                time.sleep(0.2)
            time.sleep(1.0)
        download = download_export(client, headers, job_id)
    assert status == JobStatus.SUCCEEDED.value
    assert download.status_code == 200
    assert _job(factory, job_id).attempt_count == 1


def test_queue_smoke_script_round_trips_through_a_worker(monkeypatch):
    from celery.contrib.testing.worker import start_worker

    from scripts import queue_smoke

    _purge_export_queue()
    with start_worker(
        celery_app,
        pool="solo",
        perform_ping_check=False,
        queues=[get_settings().export_queue_name],
        shutdown_timeout=30,
    ):
        assert queue_smoke.run(timeout=30) == 0
    monkeypatch.setattr(get_settings(), "celery_broker_url", "")
    assert queue_smoke.run(timeout=1) == 2
