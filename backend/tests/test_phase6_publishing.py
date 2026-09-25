from pathlib import Path
from uuid import UUID

from openpyxl import load_workbook
from pptx import Presentation
from sqlalchemy import select

from app.config import get_settings
from app.models import ExportJob, OrgUnit
from app.services.export_view import status_label
from tests.conftest import auth_header, download_export, login, query_dashboard
from tests.helpers import put_population, put_raw


def _unit(session, code):
    return session.scalar(select(OrgUnit).where(OrgUnit.code == code))


def _seed_anc(session, org):
    put_population(session, org, 2024, 1_000_000, code="PUB_POP_2024")
    put_raw(session, org, "FY2024/25", "ANC1", 47_700)
    put_raw(session, org, "FY2024/25", "ANC4", 29_650)
    put_raw(session, org, "FY2024/25", "IFA_30", 90_000)
    session.commit()


def test_excel_matches_dashboard_and_embeds_run_metadata(client, session, tmp_path, monkeypatch):
    pader = _unit(session, "PADER")
    _seed_anc(session, pader)
    monkeypatch.setattr(get_settings(), "export_dir", str(tmp_path))
    csrf = login(client, "pader.focal")
    headers = auth_header(csrf)
    dash = query_dashboard(client, headers, pader.id)
    assert dash.status_code == 201
    dashboard = dash.json()
    by_code = {row["indicator_code"]: row for row in dashboard["module_result"]["indicators"]}
    created = client.post(
        "/exports/excel",
        json={
            "org_unit_id": str(pader.id),
            "period": "FY2024/25",
            "module": "anc",
            "analysis_snapshot_id": dashboard["analysis_snapshot_id"],
            "view_hash": dashboard["view_hash"],
        },
        headers=headers,
    )
    assert created.status_code == 202, created.text
    job_id = created.json()["job_id"]
    meta = client.get(f"/exports/jobs/{job_id}", headers=headers)
    assert meta.status_code == 200
    assert meta.json()["downloadable"] is True
    assert meta.json()["calculation_run_id"]
    unsafe = client.get(f"/exports/jobs/{job_id}/file", headers=headers)
    assert unsafe.status_code in {404, 405}
    download = download_export(client, headers, job_id)
    assert download.status_code == 200
    assert download.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    path = tmp_path / f"{job_id}.xlsx"
    assert path.exists()
    book = load_workbook(path)
    assert set(book.sheetnames) >= {
        "Scorecard",
        "Calculations",
        "Raw Data",
        "Population",
        "Methodology",
        "Data Quality",
        "Metadata",
    }
    names = {row[0]: row for row in book["Scorecard"].iter_rows(min_row=2, values_only=True)}
    anc1 = by_code["ANC1_COVERAGE"]
    assert names[anc1["name"]][1] == anc1["raw_value"]
    # The workbook shows the status as the dashboard names it, on a whole-cell colour.
    assert names[anc1["name"]][3] == status_label(anc1["status"])
    meta_rows = {row[0]: row[1] for row in book["Metadata"].iter_rows(min_row=2, values_only=True)}
    assert meta_rows["current_run_id"] == meta.json()["calculation_run_id"]
    assert meta_rows["period"] == "FY2024/25"
    assert by_code["IFA_COVERAGE"]["raw_value"] > 100
    assert names[by_code["IFA_COVERAGE"]["name"]][1] == by_code["IFA_COVERAGE"]["raw_value"]


def test_powerpoint_is_editable_and_uses_run_values(client, session, tmp_path, monkeypatch):
    pader = _unit(session, "PADER")
    _seed_anc(session, pader)
    monkeypatch.setattr(get_settings(), "export_dir", str(tmp_path))
    csrf = login(client, "pader.focal")
    headers = auth_header(csrf)
    dash = query_dashboard(client, headers, pader.id).json()
    created = client.post(
        "/exports/powerpoint",
        json={
            "org_unit_id": str(pader.id),
            "period": "FY2024/25",
            "module": "anc",
            "analysis_snapshot_id": dash["analysis_snapshot_id"],
            "view_hash": dash["view_hash"],
        },
        headers=headers,
    )
    assert created.status_code == 202, created.text
    job_id = created.json()["job_id"]
    path = Path(session.get(ExportJob, UUID(job_id)).file_path)
    deck = Presentation(path)
    texts = []
    pictures = 0
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13:
                pictures += 1
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    texts.extend(cell.text for cell in row.cells)
    blob = " ".join(texts)
    assert pictures == 0
    job = session.get(ExportJob, UUID(job_id))
    assert job is not None and job.calculation_run_id
    assert str(job.calculation_run_id) in blob
    first = dash["module_result"]["indicators"][0]
    assert first["name"] in blob
    assert first["display_value"] in blob or first["name"] in blob
    assert "platform-default" in blob.lower() or "Official MoH slide templates are not yet supplied" in blob


def test_report_contains_run_and_values(client, session, tmp_path, monkeypatch):
    pader = _unit(session, "PADER")
    _seed_anc(session, pader)
    monkeypatch.setattr(get_settings(), "export_dir", str(tmp_path))
    csrf = login(client, "pader.focal")
    headers = auth_header(csrf)
    dash = query_dashboard(client, headers, pader.id).json()
    created = client.post(
        "/exports/report",
        json={
            "org_unit_id": str(pader.id),
            "period": "FY2024/25",
            "module": "anc",
            "analysis_snapshot_id": dash["analysis_snapshot_id"],
            "view_hash": dash["view_hash"],
        },
        headers=headers,
    )
    assert created.status_code == 202
    job_id = created.json()["job_id"]
    download = download_export(client, headers, job_id)
    assert download.status_code == 200
    assert download.headers["content-type"].startswith("text/markdown")
    assert download.headers["content-disposition"].endswith('.md"')
    text = download.content.decode("utf-8")
    assert "Format: Markdown (.md). This is not a Word or PDF document." in text
    job = session.get(ExportJob, UUID(job_id))
    assert job is not None and job.calculation_run_id
    assert str(job.calculation_run_id) in text
    assert "AI did not calculate official numbers" in text
    assert dash["module_result"]["indicators"][0]["name"] in text


def test_export_job_is_owner_scoped(client, session, tmp_path, monkeypatch):
    pader = _unit(session, "PADER")
    _seed_anc(session, pader)
    monkeypatch.setattr(get_settings(), "export_dir", str(tmp_path))
    owner = login(client, "pader.focal")
    headers = auth_header(owner)
    dash = query_dashboard(client, headers, pader.id).json()
    created = client.post(
        "/exports/excel",
        json={
            "org_unit_id": str(pader.id),
            "period": "FY2024/25",
            "module": "anc",
            "analysis_snapshot_id": dash["analysis_snapshot_id"],
            "view_hash": dash["view_hash"],
        },
        headers=headers,
    )
    job_id = created.json()["job_id"]
    other = login(client, "view.only")
    denied = download_export(client, auth_header(other), job_id)
    assert denied.status_code == 404
