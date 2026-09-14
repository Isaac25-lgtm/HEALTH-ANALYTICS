"""Formula versions are selected for the period's effective date (period end) and fail closed.

Dated history is authoritative. Undated legacy versions are used only when an indicator has no
dated history and the environment's governance policy permits it. Every decision, including
unavailability, is recorded in calculation-run provenance and survives into the snapshot, AI
evidence and exports.
"""

from __future__ import annotations

import json
import random
from datetime import date
from uuid import UUID

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.config import get_settings
from app.models import CalculatedValue, ExportJob, Indicator, IndicatorVersion, OrgUnit
from app.services.calculation import (
    FORMULA_VERSION_UNAVAILABLE,
    resolve_indicator_versions_for_period,
    run_calculation,
    select_indicator_versions_for_period,
)
from app.services.evidence import evidence_package
from tests.conftest import auth_header, login, query_dashboard
from tests.helpers import put_population, put_raw

CODE = "ANC1_COVERAGE"


def _indicator(session) -> Indicator:
    return session.scalar(select(Indicator).where(Indicator.code == CODE))


def _seeded(session) -> IndicatorVersion:
    indicator = _indicator(session)
    return session.scalar(
        select(IndicatorVersion).where(
            IndicatorVersion.indicator_id == indicator.id, IndicatorVersion.is_current.is_(True)
        )
    )


def _clone(session, name: str, valid_from: date | None, valid_to: date | None, current: bool = False):
    base = _seeded(session)
    row = IndicatorVersion(
        indicator_id=base.indicator_id,
        formula_version=name,
        numerator_definition=base.numerator_definition,
        denominator_type=base.denominator_type,
        denominator_coefficient=base.denominator_coefficient,
        multiplier=base.multiplier,
        unit=base.unit,
        display_precision=base.display_precision,
        direction=base.direction,
        target=base.target,
        green_band=base.green_band,
        yellow_band=base.yellow_band,
        red_band=base.red_band,
        blue_rule=base.blue_rule,
        aggregation_method=base.aggregation_method,
        period_adjustment=base.period_adjustment,
        methodology_text=base.methodology_text,
        formula_spec=base.formula_spec,
        classification_spec=base.classification_spec,
        valid_from=valid_from,
        valid_to=valid_to,
        is_current=current,
    )
    session.add(row)
    session.flush()
    return row


def _decision(session, period: str, loaded=None):
    indicator = _indicator(session)
    return resolve_indicator_versions_for_period(session, period, loaded=loaded)[indicator.id]


def test_historical_period_uses_the_version_in_force_at_period_end(session):
    old = _clone(session, "v-2020", date(2020, 7, 1), date(2025, 6, 30))
    new = _clone(session, "v-2025", date(2025, 7, 1), None)
    past = _decision(session, "FY2024/25")
    later = _decision(session, "FY2025/26")
    assert past.version.id == old.id and past.rule == "dated_in_force"
    assert past.effective_date == date(2025, 6, 30)
    assert later.version.id == new.id
    # The undated seeded version is ignored once dated history exists.
    assert _seeded(session).id not in {past.version.id, later.version.id}


def test_future_only_dated_version_makes_the_indicator_unavailable(session):
    _clone(session, "v-2030", date(2030, 7, 1), None)
    decision = _decision(session, "FY2024/25")
    assert decision.version is None
    assert decision.rule == "no_dated_version_in_force"
    assert decision.reason_code == FORMULA_VERSION_UNAVAILABLE
    indicator = _indicator(session)
    assert indicator.id not in {row.indicator_id for row in select_indicator_versions_for_period(session, "FY2024/25")}


def test_expired_dated_version_is_not_used_for_a_later_period(session):
    _clone(session, "v-expired", date(2018, 7, 1), date(2020, 6, 30))
    assert _decision(session, "FY2024/25").rule == "no_dated_version_in_force"
    assert _decision(session, "FY2019/20").version.formula_version == "v-expired"


def test_gap_between_dated_versions_is_unavailable(session):
    _clone(session, "v-a", date(2020, 7, 1), date(2024, 6, 30))
    _clone(session, "v-b", date(2025, 7, 1), None)
    gap = _decision(session, "FY2024/25")
    assert gap.version is None and gap.rule == "no_dated_version_in_force"
    assert _decision(session, "FY2023/24").version.formula_version == "v-a"
    assert _decision(session, "FY2025/26").version.formula_version == "v-b"


def test_overlapping_versions_choose_the_latest_valid_from_and_record_candidates(session):
    open_ended = _clone(session, "v-open", date(2020, 7, 1), None)
    newer = _clone(session, "v-newer", date(2024, 7, 1), None)
    decision = _decision(session, "FY2024/25")
    assert decision.version.id == newer.id
    assert decision.rule == "dated_overlap_latest_valid_from"
    assert set(decision.candidate_version_ids) == {str(open_ended.id), str(newer.id)}


def test_ties_are_deterministic_and_fail_closed_without_a_governance_marker(session):
    first = _clone(session, "v-tie-1", date(2024, 7, 1), None)
    second = _clone(session, "v-tie-2", date(2024, 7, 1), None)
    loaded = list(session.scalars(select(IndicatorVersion)).all())
    outcomes = set()
    for seed in range(6):
        random.Random(seed).shuffle(loaded)
        decision = _decision(session, "FY2024/25", loaded=list(loaded))
        outcomes.add((decision.rule, decision.version.id if decision.version else None))
    assert outcomes == {("ambiguous_dated_versions", None)}
    seeded = _seeded(session)
    seeded.is_current = False
    session.flush()
    second.is_current = True
    session.flush()
    loaded = list(session.scalars(select(IndicatorVersion)).all())
    for order in (loaded, list(reversed(loaded))):
        decision = _decision(session, "FY2024/25", loaded=order)
        assert decision.version.id == second.id
        assert decision.rule == "dated_tie_resolved_by_current_flag"
    assert first.id != second.id


def test_undated_legacy_versions_follow_the_environment_policy(session, monkeypatch):
    seeded = _seeded(session)
    decision = _decision(session, "FY2024/25")
    assert decision.version.id == seeded.id
    assert decision.rule == "undated_current_fallback"
    assert decision.policy == "development_default"

    monkeypatch.setattr(get_settings(), "formula_undated_fallback", False)
    disabled = _decision(session, "FY2024/25")
    assert disabled.version is None and disabled.rule == "undated_fallback_not_permitted"

    monkeypatch.setattr(get_settings(), "formula_undated_fallback", None)
    monkeypatch.setattr(get_settings(), "app_env", "production")
    unconfigured = _decision(session, "FY2024/25")
    assert unconfigured.version is None and unconfigured.policy == "not_configured"

    monkeypatch.setattr(get_settings(), "formula_undated_fallback", True)
    enabled = _decision(session, "FY2024/25")
    assert enabled.version.id == seeded.id and enabled.policy == "explicitly_enabled"


def test_run_provenance_records_selection_and_unavailability(session):
    uganda = session.scalar(select(OrgUnit).where(OrgUnit.code == "UG"))
    put_population(session, uganda, 2024, 10_000)
    put_raw(session, uganda, "FY2024/25", "ANC1", 500)
    put_raw(session, uganda, "FY2024/25", "ANC4", 300)
    _clone(session, "v-2030", date(2030, 7, 1), None)
    run = run_calculation(
        session,
        org_unit=uganda,
        period="FY2024/25",
        user=None,
        indicator_codes=[CODE, "ANC4_COVERAGE"],
    )
    config = run.config_snapshot
    selection = {item["indicator_code"]: item for item in config["formula_version_selection"]}
    assert selection[CODE]["status"] == "unavailable"
    assert selection[CODE]["reason_code"] == FORMULA_VERSION_UNAVAILABLE
    assert selection[CODE]["rule"] == "no_dated_version_in_force"
    assert selection[CODE]["effective_date"] == "2025-06-30"
    assert selection["ANC4_COVERAGE"]["status"] == "selected"
    assert selection["ANC4_COVERAGE"]["rule"] == "undated_current_fallback"
    assert config["unavailable_indicators"] == [
        {"indicator_code": CODE, "reason_code": FORMULA_VERSION_UNAVAILABLE, "rule": "no_dated_version_in_force"}
    ]
    assert config["formula_version_policy"] == {
        "effective_date": "2025-06-30",
        "effective_date_rule": "period_end",
        "undated_fallback_policy": "development_default",
    }
    stored = session.scalars(select(CalculatedValue).where(CalculatedValue.calculation_run_id == run.id)).all()
    codes = {session.get(IndicatorVersion, row.indicator_version_id).indicator_id for row in stored}
    assert _indicator(session).id not in codes
    versions = {item["indicator_code"]: item for item in config["indicator_versions"]}
    assert versions["ANC4_COVERAGE"]["formula_version_rule"] == "undated_current_fallback"
    assert CODE not in versions


@pytest.fixture()
def unavailable_snapshot(client, session, tmp_path, monkeypatch):
    pader = session.scalar(select(OrgUnit).where(OrgUnit.code == "PADER"))
    put_population(session, pader, 2024, 1_000_000, code="FV_POP")
    put_raw(session, pader, "FY2024/25", "ANC1", 47_700)
    put_raw(session, pader, "FY2024/25", "ANC4", 29_650)
    _clone(session, "v-2030", date(2030, 7, 1), None)
    session.commit()
    monkeypatch.setattr(get_settings(), "export_dir", str(tmp_path))
    headers = auth_header(login(client, "pader.focal"))
    dash = query_dashboard(client, headers, pader.id)
    assert dash.status_code == 201, dash.text
    return pader, headers, dash.json()


def test_unavailability_reaches_the_snapshot_and_survives_reopening(client, unavailable_snapshot):
    _pader, headers, dash = unavailable_snapshot
    rows = {row["indicator_code"]: row for row in dash["module_result"]["indicators"]}
    anc1 = rows[CODE]
    assert anc1["reason_code"] == FORMULA_VERSION_UNAVAILABLE
    assert anc1["raw_value"] is None and anc1["display_value"] is None
    assert anc1["name"] and anc1["formula_version_rule"] == "no_dated_version_in_force"
    assert rows["ANC4_COVERAGE"]["raw_value"] is not None
    reopened = client.get(f"/analysis-snapshots/{dash['analysis_snapshot_id']}", headers=headers).json()
    reopened_rows = {row["indicator_code"]: row for row in reopened["module_result"]["indicators"]}
    assert reopened_rows[CODE]["reason_code"] == FORMULA_VERSION_UNAVAILABLE


def test_unavailability_is_preserved_in_ai_evidence_and_answers(client, unavailable_snapshot):
    pader, headers, dash = unavailable_snapshot
    package = evidence_package(dash)
    evidence = {row["indicator_code"]: row for row in package["indicators"]}
    assert evidence[CODE]["reason_code"] == FORMULA_VERSION_UNAVAILABLE
    assert evidence[CODE]["formula_version_rule"] == "no_dated_version_in_force"
    assert "raw_value" not in evidence[CODE]
    explained = client.post(
        "/ai/explain",
        json={
            "org_unit_id": str(pader.id),
            "period": "FY2024/25",
            "module": "anc",
            "indicator_code": CODE,
            "analysis_snapshot_id": dash["analysis_snapshot_id"],
            "view_hash": dash["view_hash"],
        },
        headers=headers,
    )
    assert explained.status_code == 200, explained.text
    assert "No formula version is valid" in json.dumps(explained.json())


def test_unavailability_is_preserved_in_exports(client, session, unavailable_snapshot, tmp_path):
    pader, headers, dash = unavailable_snapshot
    body = {
        "org_unit_id": str(pader.id),
        "period": "FY2024/25",
        "module": "anc",
        "analysis_snapshot_id": dash["analysis_snapshot_id"],
        "view_hash": dash["view_hash"],
    }
    excel = client.post("/exports/excel", json=body, headers=headers)
    report = client.post("/exports/report", json=body, headers=headers)
    assert excel.status_code == 202 and report.status_code == 202
    book = load_workbook(tmp_path / f"{excel.json()['job_id']}.xlsx")
    calc = list(book["Calculations"].iter_rows(values_only=True))
    header = calc[0]
    row = next(item for item in calc[1:] if item[0] == CODE)
    assert row[header.index("Reason code")] == FORMULA_VERSION_UNAVAILABLE
    assert row[header.index("Formula version rule")] == "no_dated_version_in_force"
    assert row[header.index("Exact value")] is None
    job = session.get(ExportJob, UUID(report.json()["job_id"]))
    text = open(job.file_path, encoding="utf-8").read()
    assert FORMULA_VERSION_UNAVAILABLE in text
