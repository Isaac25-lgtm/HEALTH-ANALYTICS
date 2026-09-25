"""Colour-coded Excel workbook written from a committed snapshot.

Cells hold the exact stored values; number formats round them for display the way the dashboard
does, so a reader sees "92.4%" while the workbook keeps the full value for traceability.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.export_view import (
    NAVY,
    NAVY_LIGHT,
    STATUS_FILL,
    STATUS_INK,
    ExportView,
    change_text,
    excel_number_format,
    interpretation_text,
    resolve_status,
    status_label,
    thresholds_text,
)
from app.version import SOFTWARE_VERSION

WHITE = "FFFFFF"
_THIN = Side(style="thin", color="FFFFFF")
_GRID = Side(style="thin", color="D5DEE8")
CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
GRID_BORDER = Border(left=_GRID, right=_GRID, top=_GRID, bottom=_GRID)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color=WHITE, size=10)
TITLE_FONT = Font(bold=True, color=NAVY, size=16)
SUBTITLE_FONT = Font(color="475569", size=10)
NOT_RECORDED = "Not recorded in the governed evidence"


def _fill(status: str) -> PatternFill:
    return PatternFill("solid", fgColor=STATUS_FILL.get(status, STATUS_FILL["n_a"]))


def _ink(status: str, bold: bool = True) -> Font:
    return Font(bold=bold, color=STATUS_INK.get(status, STATUS_INK["n_a"]), size=10)


def _title(sheet: Worksheet, view: ExportView, heading: str, width: int) -> int:
    """Title, subtitle and legend. Returns the header row number."""
    sheet["A1"] = f"{heading} — {view.scope_name}"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        f"{view.module_title} · {view.subtitle} · generated {view.generated_at} · "
        "values from the committed calculation run"
    )
    sheet["A2"].font = SUBTITLE_FONT
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(width, 2))
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(width, 2))
    sheet["A3"] = "Legend"
    sheet["A3"].font = Font(bold=True, size=9, color="475569")
    for offset, status in enumerate(("green", "yellow", "red", "blue", "missing"), start=2):
        cell = sheet.cell(row=3, column=offset, value=status_label(status))
        cell.fill = _fill(status)
        cell.font = Font(bold=True, size=9, color=STATUS_INK[status])
        cell.alignment = Alignment(horizontal="center")
    sheet.row_dimensions[1].height = 24
    return 5


def _header(sheet: Worksheet, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = sheet.cell(row=row, column=column, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER
    sheet.row_dimensions[row].height = 32


def _widths(sheet: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _value_cell(sheet: Worksheet, row: int, column: int, measure: dict | None) -> None:
    """Exact value, dashboard rounding via number format, whole cell coloured by status."""
    status = resolve_status(measure)
    cell = sheet.cell(row=row, column=column)
    raw = (measure or {}).get("raw_value")
    if raw is None:
        cell.value = "No data"
    else:
        cell.value = raw
        cell.number_format = excel_number_format((measure or {}).get("unit"), (measure or {}).get("display_precision"))
    cell.fill = _fill(status)
    cell.font = _ink(status)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = CELL_BORDER


def _plain_number(cell, value) -> None:
    cell.value = value
    if isinstance(value, (int, float)):
        cell.number_format = "#,##0" if float(value).is_integer() else "#,##0.0"
    cell.alignment = Alignment(horizontal="right")


def _print_setup(sheet: Worksheet) -> None:
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_view.showGridLines = False


def _scorecard(sheet: Worksheet, view: ExportView) -> None:
    labels = [
        "Indicator",
        "Value",
        "Unit",
        "Status",
        f"Comparison ({view.comparison_label or 'none'})",
        "Change",
        "Interpretation",
        "Approved bands",
        "Numerator",
        "Denominator",
    ]
    header_row = _title(sheet, view, "Indicator scorecard", len(labels))
    _header(sheet, header_row, labels)
    row_number = header_row
    for indicator in view.indicators:
        row_number += 1
        status = resolve_status(indicator)
        name = sheet.cell(row=row_number, column=1, value=indicator.get("name") or indicator.get("indicator_code"))
        name.font = Font(bold=True, color="1E293B", size=10)
        name.border = GRID_BORDER
        _value_cell(sheet, row_number, 2, indicator)
        sheet.cell(row=row_number, column=3, value=indicator.get("unit"))
        label = sheet.cell(row=row_number, column=4, value=status_label(status))
        label.fill = _fill(status)
        label.font = _ink(status)
        label.alignment = Alignment(horizontal="center")
        label.border = CELL_BORDER
        comparison = indicator.get("comparison") or {}
        if comparison.get("raw_value") is not None:
            _value_cell(sheet, row_number, 5, comparison)
        else:
            sheet.cell(row=row_number, column=5, value="No data").alignment = Alignment(horizontal="center")
        sheet.cell(row=row_number, column=6, value=change_text(indicator)).alignment = Alignment(horizontal="center")
        sheet.cell(row=row_number, column=7, value=interpretation_text(indicator) or None)
        sheet.cell(row=row_number, column=8, value=thresholds_text(indicator) or None).alignment = Alignment(
            wrap_text=True, vertical="center"
        )
        _plain_number(sheet.cell(row=row_number, column=9), indicator.get("numerator"))
        _plain_number(sheet.cell(row=row_number, column=10), indicator.get("denominator"))
        sheet.row_dimensions[row_number].height = 32 if len(thresholds_text(indicator)) > 60 else 20
    _widths(sheet, [38, 13, 15, 20, 18, 14, 16, 50, 14, 14])
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=2)
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(labels))}{max(row_number, header_row)}"
    _print_setup(sheet)


def _units(sheet: Worksheet, view: ExportView) -> None:
    columns = view.unit_columns
    labels = [view.unit_label[:-1] if view.unit_label.endswith("s") else view.unit_label] + [
        str(item.get("name") or item.get("indicator_code")) for item in columns
    ]
    labels[0] = "Unit"
    header_row = _title(sheet, view, f"{view.unit_label} scorecard", len(labels))
    _header(sheet, header_row, labels)
    row_number = header_row
    for unit in view.units:
        row_number += 1
        name = sheet.cell(row=row_number, column=1, value=unit.get("org_unit_name"))
        name.font = Font(bold=True, color="1E293B", size=10)
        name.border = GRID_BORDER
        values = unit.get("values") or {}
        for offset, column in enumerate(columns, start=2):
            _value_cell(sheet, row_number, offset, values.get(column.get("indicator_code")))
    if not view.units:
        sheet.cell(row=header_row + 1, column=1, value="No unit comparison is available for this scope.")
    _widths(sheet, [30] + [16] * len(columns))
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=2)
    if view.units:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(labels))}{row_number}"
    _print_setup(sheet)


def _ranking(sheet: Worksheet, view: ExportView) -> None:
    ranking = view.ranking
    header_row = _title(sheet, view, f"Top and bottom performers · {view.ranking_indicator_name}", 9)
    if not ranking.get("ranking_allowed"):
        sheet.cell(
            row=header_row, column=1, value=ranking.get("reason") or "Ranking is not available for this indicator."
        )
        _widths(sheet, [30, 14, 18])
        return
    for start, title, rows in (
        (1, "Best performing", ranking.get("best") or []),
        (6, "Needs attention", ranking.get("worst") or []),
    ):
        sheet.cell(row=header_row - 1, column=start, value=title).font = Font(bold=True, color=NAVY, size=11)
        for column, label in enumerate(["Rank", "Unit", "Value", "Status"], start=start):
            cell = sheet.cell(row=header_row, column=column, value=label)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")
        for index, entry in enumerate(rows, start=1):
            row_number = header_row + index
            status = resolve_status(entry)
            sheet.cell(row=row_number, column=start, value=index).alignment = Alignment(horizontal="center")
            sheet.cell(row=row_number, column=start + 1, value=entry.get("org_unit_name") or entry.get("org_unit_code"))
            _value_cell(sheet, row_number, start + 2, entry)
            label = sheet.cell(row=row_number, column=start + 3, value=status_label(status))
            label.fill = _fill(status)
            label.font = _ink(status)
    _widths(sheet, [7, 28, 12, 18, 3, 7, 28, 12, 18])
    _print_setup(sheet)


def _trends(sheet: Worksheet, view: ExportView) -> None:
    columns = view.indicators
    labels = ["Period"] + [str(item.get("name") or item.get("indicator_code")) for item in columns]
    header_row = _title(sheet, view, "Trends", min(len(labels), 12))
    _header(sheet, header_row, labels)
    row_number = header_row
    for trend in view.trends:
        row_number += 1
        sheet.cell(row=row_number, column=1, value=str(trend.get("period"))).font = Font(bold=True, size=10)
        values = trend.get("values") or {}
        for offset, column in enumerate(columns, start=2):
            _value_cell(sheet, row_number, offset, values.get(column.get("indicator_code")))
    _widths(sheet, [14] + [14] * len(columns))
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=2)
    if view.trends and view.trend_indicator:
        code = view.trend_indicator.get("indicator_code")
        position = next(
            (index for index, item in enumerate(columns, start=2) if item.get("indicator_code") == code), None
        )
        if position is not None:
            chart = LineChart()
            chart.title = f"{view.trend_indicator.get('name')} — trend"
            chart.height = 8
            chart.width = 22
            chart.legend = None
            data = Reference(sheet, min_col=position, min_row=header_row, max_row=row_number)
            categories = Reference(sheet, min_col=1, min_row=header_row + 1, max_row=row_number)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            series = chart.series[0]
            series.graphicalProperties.line.solidFill = NAVY
            series.graphicalProperties.line.width = 28000
            series.smooth = False
            sheet.add_chart(chart, f"B{row_number + 3}")


def _style_table(sheet: Worksheet, widths: list[int]) -> None:
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    for index, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        if index % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=NAVY_LIGHT)
    _widths(sheet, widths)
    sheet.freeze_panes = "A2"


def _governed(row: dict, fallback: dict, key: str) -> object:
    """Governed evidence first, then the committed snapshot row; never a silent blank."""
    value = row.get(key)
    if value in (None, ""):
        value = fallback.get(key)
    if isinstance(value, dict):
        value = value.get("text") or value.get("state")
    return value if value not in (None, "") else NOT_RECORDED


def _freshness_text(package: dict, dashboard: dict, row: dict, fallback: dict) -> str:
    for candidate in (row.get("source_freshness_at"), fallback.get("source_freshness_at")):
        if candidate:
            return str(candidate)
    freshness = package.get("freshness") or (dashboard.get("module_result") or {}).get("freshness") or {}
    return str(freshness.get("latest_source_freshness_at") or freshness.get("latest_extracted_at") or NOT_RECORDED)


def write_excel(view: ExportView, dashboard: dict, package: dict, destination: Path, template_version: str) -> None:
    book = Workbook()
    _scorecard(book.active, view)
    book.active.title = "Scorecard"
    _units(book.create_sheet("Units"), view)
    _ranking(book.create_sheet("Ranking"), view)
    _trends(book.create_sheet("Trends"), view)

    snapshot_rows = {
        row.get("indicator_code"): row
        for row in (dashboard.get("module_result") or {}).get("indicators") or []
        if row.get("indicator_code")
    }
    calc = book.create_sheet("Calculations")
    calc.append(
        [
            "Code",
            "Display",
            "Exact value",
            "Numerator",
            "Denominator",
            "Formula version",
            "Thresholds",
            "Status rule",
            "Change interpretation",
            "Run",
            "BLUE / unavailable reason",
            "Reason code",
            "Formula version rule",
        ]
    )
    for row in package["indicators"]:
        fallback = snapshot_rows.get(row.get("indicator_code"), {})
        calc.append(
            [
                row.get("indicator_code"),
                row.get("display_value"),
                row.get("raw_value"),
                row.get("numerator"),
                row.get("denominator"),
                _governed(row, fallback, "formula_version"),
                _governed(row, fallback, "thresholds"),
                row.get("status"),
                (row.get("change") or {}).get("interpretation") or row.get("interpretation") or "not_interpreted",
                row.get("calculation_run_id"),
                row.get("blue_reason"),
                row.get("reason_code"),
                _governed(row, fallback, "formula_version_rule"),
            ]
        )
    _style_table(calc, [26, 10, 14, 14, 14, 16, 44, 12, 18, 38, 30, 24, 26])

    raw = book.create_sheet("Raw Data")
    raw.append(
        [
            "Code",
            "Numerator",
            "Denominator",
            "Mapping version",
            "Source freshness",
            "Aggregation scope",
            "Identifiers",
            "Missing is not converted to zero",
        ]
    )
    for row in package["indicators"]:
        fallback = snapshot_rows.get(row.get("indicator_code"), {})
        raw.append(
            [
                row.get("indicator_code"),
                row.get("numerator"),
                row.get("denominator"),
                _governed(row, fallback, "mapping_version"),
                _freshness_text(package, dashboard, row, fallback),
                _governed(row, fallback, "aggregation_policy"),
                "redacted",
                "retained",
            ]
        )
    _style_table(raw, [26, 14, 14, 18, 30, 30, 12, 16])

    pop = book.create_sheet("Population")
    population = package.get("population") or {}
    pop.append(
        ["Year", "Value", "Source", "Status", "Reason", "Version code", "Version ID", "Entry ID", "Type", "Approval"]
    )
    pop.append(
        [
            population.get("year"),
            population.get("population"),
            population.get("source"),
            population.get("status"),
            population.get("reason"),
            population.get("version_code"),
            population.get("version_id") or population.get("population_version_id"),
            population.get("used_facility_entry_id") or population.get("facility_population_entry_id"),
            population.get("population_type"),
            population.get("approval_status"),
        ]
    )
    _style_table(pop, [8, 14, 60, 12, 30, 24, 38, 38, 14, 12])
    pop["B2"].number_format = "#,##0"

    method = book.create_sheet("Methodology")
    method.append(["Item", "Detail"])
    method.append(["Template", template_version])
    method.append(["Software", SOFTWARE_VERSION])
    method.append(["Values", "Exact stored values; cell formats round for display only."])
    method.append(["Colours", "Whole cells are coloured by the approved status of each value."])
    method.append(["Note", "Formulas remain in the versioned indicator catalogue. AI did not calculate these values."])
    _style_table(method, [16, 90])

    quality = book.create_sheet("Data Quality")
    quality.append(["Rule", "Severity", "Explanation", "Occurrences"])
    for flag in package.get("quality_flags") or []:
        quality.append(
            [flag.get("rule_id"), flag.get("severity"), flag.get("explanation"), flag.get("occurrences") or 1]
        )
    _style_table(quality, [32, 12, 90, 12])

    meta = book.create_sheet("Metadata")
    meta.append(["Field", "Value"])
    for key in ("period", "comparison_period", "module", "current_run_id", "analysis_snapshot_id", "view_hash"):
        meta.append([key, package.get(key) or dashboard.get(key)])
    meta.append(["scope", (package.get("scope") or {}).get("org_unit_name")])
    meta.append(["indicator_count", len(package["indicators"])])
    meta.append(["generated_at", view.generated_at])
    _style_table(meta, [22, 60])
    book.save(destination)
