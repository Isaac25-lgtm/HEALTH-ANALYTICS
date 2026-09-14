"""Owner-supplied district/city population workbook: verify, reconcile, then import as draft.

Rules (amendment §2/§3 and the owner's population instructions):

* The workbook is never modified. Its SHA-256 is verified before every read and every
  import; a different checksum is a different source that needs a new version decision.
* The NATIONAL TOTAL row is never a district or city. It is excluded by default and may
  only be imported as a direct country value when explicitly requested.
* Units match internal organisation units only by exact normalised name plus unit type,
  or through an approved, audited alias. There is no fuzzy matching.
* Unmatched, ambiguous, pending, rejected, type-mismatched and duplicate-target rows block
  the import. The Region column is descriptive metadata, never a parent.
* Imports create DRAFT versions that preserve source name, type, region, year, value,
  workbook identity and import date. Facilities and sub-counties never receive values.
"""

from __future__ import annotations

import hashlib
from collections import Counter
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from uuid import UUID

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.enums import (
    ActionPermission,
    AggregationClass,
    AliasDecisionStatus,
    ApprovalStatus,
    OrgUnitLevel,
    PopulationType,
    aggregation_class,
)
from app.models import (
    OrgUnit,
    PopulationImportBatch,
    PopulationImportRow,
    PopulationSourceAlias,
    PopulationValue,
    PopulationVersion,
    User,
)
from app.services.audit import write_audit
from app.services.authorization import (
    AuthorizationError,
    authorised_org_unit_ids,
    require_action,
    require_org_unit_access,
)

VERIFIED_SHA256 = "5AE43DCA4533347FBB95C3F9B4E08A25C953C90905445614BA7D4E4BD3E75072"
SOURCE_DATASET = "UBOS_NPHC2024_DISTRICT_CITY_PROJECTIONS_2025_2030"
SOURCE_NAME = "UBOS NPHC 2024 final census counts and UBOS 2025-2030 district/city mid-year projections"
SHEET_NAME = "District_City_Populations"
HEADER = (
    "Administrative Unit",
    "Type",
    "Region",
    "2024 Census",
    "2025 Projection",
    "2026 Projection",
    "2027 Projection",
    "2028 Projection",
    "2029 Projection",
    "2030 Projection",
)
NATIONAL_TOTAL_LABEL = "NATIONAL TOTAL"
TYPE_LEVELS = {"District": OrgUnitLevel.DISTRICT.value, "City": OrgUnitLevel.CITY.value}
CENSUS_YEAR = 2024
# Structure the owner-approved source is expected to have. Differences are reported, never
# silently corrected.
EXPECTED_DISTRICTS = 135
EXPECTED_CITIES = 11
EXPECTED_UNITS = EXPECTED_DISTRICTS + EXPECTED_CITIES
EXPECTED_YEARS = list(range(2024, 2031))
# The four broad statistical regions in the workbook. These are descriptive metadata for
# reconciliation only; they are never an analytical parent and never a health sub-region.
BROAD_REGIONS = ("Central", "Eastern", "Northern", "Western")
IMPORTER_VERSION = "hpip-population-workbook-2"
SOURCE_TYPE = "national_statistical_office_workbook"
# Candidate matches are only as authoritative as the organisation units they matched against.
REFERENCE_SYNTHETIC = "synthetic_development_fixtures"
REFERENCE_AUTHORITATIVE = "authoritative_org_units"
STAGED = "staged"
REVIEW_PENDING = "pending_review"

LOWER_LEVEL_NOTE = (
    "Facility and sub-county populations are never derived from this workbook. "
    "They remain unavailable until separately supplied and approved."
)

BLOCKING_STATUSES = {
    "unmatched",
    "ambiguous",
    "type_mismatch",
    "pending_alias",
    "rejected_alias",
    "invalid_alias",
    "duplicate_target",
}


class PopulationWorkbookError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class WorkbookUnit:
    row_number: int
    name: str
    unit_type: str
    region: str | None
    values: dict[int, int]


@dataclass
class WorkbookExtract:
    file_name: str
    sha256: str
    sheet: str
    years: list[int]
    column_labels: dict[int, str]
    units: list[WorkbookUnit]
    national_total: dict[int, int]

    @property
    def cell_count(self) -> int:
        return sum(len(unit.values) for unit in self.units)

    def summary(self) -> dict:
        types = Counter(unit.unit_type for unit in self.units)
        return {
            "file_name": self.file_name,
            "sha256": self.sha256,
            "sheet": self.sheet,
            "unit_rows": len(self.units),
            "districts": types.get("District", 0),
            "cities": types.get("City", 0),
            "years": self.years,
            "cells": self.cell_count,
            "national_total": {
                "row_label": NATIONAL_TOTAL_LABEL,
                "values": {str(year): value for year, value in self.national_total.items()},
                "equals_unit_sum": True,
            },
        }


@dataclass
class UnitReconciliation:
    unit: WorkbookUnit
    status: str
    org_unit: OrgUnit | None = None
    alias: PopulationSourceAlias | None = None
    detail: str | None = None

    def as_dict(self) -> dict:
        return {
            "row_number": self.unit.row_number,
            "source_unit_name": self.unit.name,
            "source_unit_type": self.unit.unit_type,
            "source_region": self.unit.region,
            "status": self.status,
            "org_unit_id": str(self.org_unit.id) if self.org_unit else None,
            "org_unit_code": self.org_unit.code if self.org_unit else None,
            "org_unit_name": self.org_unit.name if self.org_unit else None,
            "org_unit_level": self.org_unit.level_type if self.org_unit else None,
            "alias_id": str(self.alias.id) if self.alias else None,
            "detail": self.detail,
        }


@dataclass
class ReconciliationReport:
    extract: WorkbookExtract
    units: list[UnitReconciliation]
    internal_units_without_source: list[OrgUnit] = field(default_factory=list)

    @property
    def counts(self) -> dict[str, int]:
        return dict(sorted(Counter(item.status for item in self.units).items()))

    @property
    def blocking(self) -> list[UnitReconciliation]:
        return [item for item in self.units if item.status in BLOCKING_STATUSES]

    @property
    def can_apply(self) -> bool:
        return not self.blocking

    def as_dict(self, *, mode: str = "dry_run") -> dict:
        return {
            "mode": mode,
            "source_dataset": SOURCE_DATASET,
            "workbook": {**self.extract.summary(), "checksum_verified": True},
            "counts": self.counts,
            "can_apply": self.can_apply,
            "blocking_issues": [item.as_dict() for item in self.blocking],
            "units": [item.as_dict() for item in self.units],
            "internal_units_without_source": [
                {"id": str(unit.id), "code": unit.code, "name": unit.name, "level_type": unit.level_type}
                for unit in self.internal_units_without_source
            ],
            "national_total_handling": "excluded_unless_explicitly_imported_as_a_direct_country_value",
            "lower_levels": LOWER_LEVEL_NOTE,
        }


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def verify_checksum(path: Path, expected_sha256: str = VERIFIED_SHA256) -> str:
    if not path.is_file():
        raise PopulationWorkbookError("workbook_missing", f"Population workbook not found: {path.name}")
    actual = file_sha256(path)
    if actual != expected_sha256.upper():
        raise PopulationWorkbookError(
            "checksum_mismatch",
            "The workbook checksum differs from the verified source. A changed file is a new source "
            "and requires a new source/version decision before it can be read or imported.",
        )
    return actual


def normalise_name(value: object) -> str:
    return " ".join(str(value or "").split()).casefold()


def _positive_int(value: object, *, row_number: int, label: str) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise PopulationWorkbookError("invalid_value", f"Row {row_number} {label} is not numeric.")
    if float(value) <= 0 or float(value) != int(value):
        raise PopulationWorkbookError("invalid_value", f"Row {row_number} {label} must be a positive whole number.")
    return int(value)


def read_population_workbook(path: Path, *, expected_sha256: str = VERIFIED_SHA256) -> WorkbookExtract:
    """Read the workbook without modifying it. The checksum is verified before and after reading."""
    from openpyxl import load_workbook

    sha = verify_checksum(path, expected_sha256)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in book.sheetnames:
            raise PopulationWorkbookError("sheet_missing", f"Sheet {SHEET_NAME} is missing.")
        rows = list(book[SHEET_NAME].iter_rows(values_only=True))
    finally:
        book.close()
    header_index = next(
        (index for index, row in enumerate(rows) if row and tuple(row[: len(HEADER)]) == HEADER),
        None,
    )
    if header_index is None:
        raise PopulationWorkbookError("header_mismatch", "The expected population header row was not found.")
    years = [int(label[:4]) for label in HEADER[3:]]
    labels = dict(zip(years, HEADER[3:], strict=True))
    units: list[WorkbookUnit] = []
    national: dict[int, int] = {}
    seen: Counter[str] = Counter()
    for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if not row or all(cell is None for cell in row):
            continue
        name = " ".join(str(row[0] or "").split())
        unit_type = row[1]
        values = {
            year: _positive_int(row[3 + index], row_number=offset, label=labels[year])
            for index, year in enumerate(years)
        }
        if name.upper() == NATIONAL_TOTAL_LABEL and unit_type is None:
            if national:
                raise PopulationWorkbookError("duplicate_national_total", "More than one NATIONAL TOTAL row exists.")
            national = values
            continue
        if unit_type not in TYPE_LEVELS:
            raise PopulationWorkbookError("invalid_unit_type", f"Row {offset} has unsupported unit type {unit_type!r}.")
        if not name:
            raise PopulationWorkbookError("invalid_unit_name", f"Row {offset} has no administrative unit name.")
        seen[normalise_name(name)] += 1
        units.append(
            WorkbookUnit(
                row_number=offset,
                name=name,
                unit_type=str(unit_type),
                region=" ".join(str(row[2]).split()) if row[2] is not None else None,
                values=values,
            )
        )
    duplicates = sorted(name for name, count in seen.items() if count > 1)
    if duplicates:
        raise PopulationWorkbookError("duplicate_source_units", f"Duplicate unit names: {', '.join(duplicates)}")
    if not national:
        raise PopulationWorkbookError("national_total_missing", "The NATIONAL TOTAL row was not found.")
    for year in years:
        if sum(unit.values[year] for unit in units) != national[year]:
            raise PopulationWorkbookError(
                "national_total_mismatch", f"NATIONAL TOTAL for {year} does not equal the sum of units."
            )
    if file_sha256(path) != sha:
        raise PopulationWorkbookError("workbook_changed", "The workbook changed while it was being read.")
    return WorkbookExtract(
        file_name=path.name,
        sha256=sha,
        sheet=SHEET_NAME,
        years=years,
        column_labels=labels,
        units=units,
        national_total=national,
    )


def _candidate_units(session: Session) -> list[OrgUnit]:
    return list(
        session.scalars(
            select(OrgUnit).where(
                OrgUnit.active.is_(True),
                OrgUnit.level_type.in_(list(TYPE_LEVELS.values())),
            )
        ).all()
    )


def reconcile_workbook(session: Session, extract: WorkbookExtract) -> ReconciliationReport:
    candidates = _candidate_units(session)
    by_name: dict[str, list[OrgUnit]] = {}
    for unit in candidates:
        by_name.setdefault(normalise_name(unit.name), []).append(unit)
    aliases = {
        row.source_unit_name: row
        for row in session.scalars(
            select(PopulationSourceAlias).where(PopulationSourceAlias.source_dataset == SOURCE_DATASET)
        ).all()
    }
    results: list[UnitReconciliation] = []
    for unit in extract.units:
        results.append(_reconcile_unit(session, unit, aliases.get(unit.name), by_name))
    targets = Counter(str(item.org_unit.id) for item in results if item.org_unit is not None)
    for item in results:
        if item.org_unit is not None and targets[str(item.org_unit.id)] > 1:
            item.status = "duplicate_target"
            item.detail = "More than one workbook unit resolves to the same organisation unit."
    matched = {str(item.org_unit.id) for item in results if item.org_unit is not None}
    uncovered = sorted(
        (unit for unit in candidates if str(unit.id) not in matched), key=lambda unit: (unit.level_type, unit.name)
    )
    return ReconciliationReport(extract=extract, units=results, internal_units_without_source=uncovered)


def _reconcile_unit(
    session: Session,
    unit: WorkbookUnit,
    alias: PopulationSourceAlias | None,
    by_name: dict[str, list[OrgUnit]],
) -> UnitReconciliation:
    level = TYPE_LEVELS[unit.unit_type]
    if alias is not None:
        if alias.decision_status == AliasDecisionStatus.PROPOSED.value:
            return UnitReconciliation(
                unit, "pending_alias", alias=alias, detail="The alias awaits a recorded decision."
            )
        if alias.decision_status == AliasDecisionStatus.REJECTED.value:
            return UnitReconciliation(
                unit, "rejected_alias", alias=alias, detail="The proposed alias was rejected; a new decision is needed."
            )
        target = session.get(OrgUnit, alias.org_unit_id) if alias.org_unit_id else None
        if target is None or not target.active or not _alias_level_compatible(level, target.level_type):
            return UnitReconciliation(
                unit,
                "invalid_alias",
                alias=alias,
                detail="The approved alias target is missing, inactive or not a district-equivalent peer.",
            )
        return UnitReconciliation(unit, "matched_alias", org_unit=target, alias=alias)
    same_name = by_name.get(normalise_name(unit.name), [])
    exact = [candidate for candidate in same_name if candidate.level_type == level]
    if len(exact) == 1:
        return UnitReconciliation(unit, "matched_exact", org_unit=exact[0])
    if len(exact) > 1:
        return UnitReconciliation(unit, "ambiguous", detail="Several organisation units share this name and type.")
    if same_name:
        return UnitReconciliation(
            unit,
            "type_mismatch",
            detail="An organisation unit has this name but a different type; an explicit alias decision is required.",
        )
    return UnitReconciliation(unit, "unmatched", detail="No organisation unit has this exact name and type.")


def _alias_level_compatible(source_level: str, target_level: str) -> bool:
    return aggregation_class(source_level) == aggregation_class(target_level) == AggregationClass.DISTRICT_EQUIVALENT


def propose_alias(
    session: Session,
    user: User,
    *,
    source_unit_name: str,
    source_unit_type: str,
    org_unit_id: UUID,
    note: str,
    evidence: dict | None = None,
) -> PopulationSourceAlias:
    require_action(session, user, ActionPermission.EDIT_POPULATION)
    if source_unit_type not in TYPE_LEVELS:
        raise AuthorizationError("invalid_input", "Alias source type must be District or City.")
    if not note or not note.strip():
        raise AuthorizationError("invalid_input", "An alias proposal needs a reviewable note.")
    target = session.get(OrgUnit, org_unit_id)
    if target is None or not target.active:
        raise AuthorizationError("invalid_input", "Alias target organisation unit is not active.")
    require_org_unit_access(session, user, target.id)
    if not _alias_level_compatible(TYPE_LEVELS[source_unit_type], target.level_type):
        raise AuthorizationError("invalid_input", "District and city names may only alias district-equivalent units.")
    existing = session.scalar(
        select(PopulationSourceAlias).where(
            PopulationSourceAlias.source_dataset == SOURCE_DATASET,
            PopulationSourceAlias.source_unit_name == source_unit_name,
        )
    )
    if existing is not None and existing.decision_status != AliasDecisionStatus.REJECTED.value:
        raise AuthorizationError("invalid_input", "An alias for this source name is already proposed or approved.")
    alias = existing or PopulationSourceAlias(source_dataset=SOURCE_DATASET, source_unit_name=source_unit_name)
    alias.source_unit_type = source_unit_type
    alias.org_unit_id = target.id
    alias.target_name = target.name
    alias.decision_status = AliasDecisionStatus.PROPOSED.value
    alias.decision_note = note.strip()
    alias.proposed_by_user_id = user.id
    alias.decided_by_user_id = None
    alias.decided_at = None
    alias.evidence = evidence
    session.add(alias)
    session.flush()
    write_audit(
        session,
        actor_user_id=user.id,
        action="population_alias_proposed",
        resource_type="population_source_alias",
        resource_id=str(alias.id),
        after={"source_unit_name": source_unit_name, "target_name": target.name, "target_code": target.code},
        reason=note,
    )
    return alias


def decide_alias(
    session: Session,
    user: User,
    alias_id: UUID,
    *,
    approve: bool,
    note: str,
) -> PopulationSourceAlias:
    require_action(session, user, ActionPermission.APPROVE_POPULATION)
    alias = session.get(PopulationSourceAlias, alias_id)
    if alias is None:
        raise AuthorizationError("not_found", "Population alias not found.")
    if alias.decision_status != AliasDecisionStatus.PROPOSED.value:
        raise AuthorizationError("invalid_input", "Only proposed aliases can be decided.")
    if alias.proposed_by_user_id == user.id and not user.is_system_admin:
        raise AuthorizationError("forbidden_action", "The proposing user cannot decide this alias.")
    if not note or not note.strip():
        raise AuthorizationError("invalid_input", "An alias decision needs a recorded reason.")
    alias.decision_status = (AliasDecisionStatus.APPROVED if approve else AliasDecisionStatus.REJECTED).value
    alias.decided_by_user_id = user.id
    alias.decided_at = datetime.now(UTC)
    alias.decision_note = f"{alias.decision_note or ''}\nDecision: {note.strip()}".strip()
    write_audit(
        session,
        actor_user_id=user.id,
        action="population_alias_approved" if approve else "population_alias_rejected",
        resource_type="population_source_alias",
        resource_id=str(alias.id),
        after={"source_unit_name": alias.source_unit_name, "target_name": alias.target_name},
        reason=note,
    )
    session.flush()
    return alias


def apply_population_workbook(
    session: Session,
    user: User,
    path: Path,
    *,
    version_code: str,
    expected_sha256: str = VERIFIED_SHA256,
    include_national_total: bool = False,
    national_org_unit_code: str | None = None,
) -> list[PopulationVersion]:
    """Create DRAFT census and projection versions after a clean reconciliation.

    Approval is a separate step through ``approve_population_version``.
    """
    require_action(session, user, ActionPermission.EDIT_POPULATION)
    extract = read_population_workbook(path, expected_sha256=expected_sha256)
    report = reconcile_workbook(session, extract)
    if not report.can_apply:
        raise PopulationWorkbookError(
            "reconciliation_blocked",
            f"{len(report.blocking)} workbook unit(s) are unmatched, ambiguous or awaiting an alias decision.",
        )
    allowed = authorised_org_unit_ids(session, user)
    if any(item.org_unit.id not in allowed for item in report.units):
        raise AuthorizationError(
            "forbidden_geography", "The importing user is not authorised for every matched organisation unit."
        )
    already = session.scalar(
        select(PopulationVersion.id).where(
            PopulationVersion.source_dataset == SOURCE_DATASET,
            PopulationVersion.source_sha256 == extract.sha256,
        )
    )
    if already is not None:
        raise PopulationWorkbookError("already_imported", "This workbook checksum has already been imported.")
    code = version_code.strip()
    if not code or len(code) > 64:
        raise PopulationWorkbookError("invalid_version_code", "Version code is required (at most 64 characters).")
    national_unit = None
    if include_national_total:
        national_unit = session.scalar(select(OrgUnit).where(OrgUnit.code == (national_org_unit_code or "")))
        if national_unit is None or national_unit.level_type != OrgUnitLevel.COUNTRY.value:
            raise PopulationWorkbookError(
                "national_unit_required", "Importing the national total requires an explicit country organisation unit."
            )
        if national_unit.id not in allowed:
            raise AuthorizationError("forbidden_geography", "The importing user is not authorised nationally.")
    groups = (
        (f"{code}_CENSUS{CENSUS_YEAR}", PopulationType.CENSUS.value, [CENSUS_YEAR]),
        (
            f"{code}_PROJ{extract.years[1]}_{extract.years[-1]}",
            PopulationType.PROJECTION.value,
            [year for year in extract.years if year != CENSUS_YEAR],
        ),
    )
    imported_at = datetime.now(UTC)
    versions: list[PopulationVersion] = []
    for group_code, population_type, years in groups:
        if session.scalar(select(PopulationVersion.id).where(PopulationVersion.code == group_code)) is not None:
            raise PopulationWorkbookError("version_code_exists", f"Population version {group_code} already exists.")
        version = PopulationVersion(
            code=group_code,
            name=f"District/city {population_type} populations {years[0]}-{years[-1]}",
            source_name=SOURCE_NAME,
            source_document=extract.file_name,
            population_type=population_type,
            approval_status=ApprovalStatus.DRAFT.value,
            notes=(
                f"Imported {imported_at.isoformat()} from sheet {extract.sheet}. Region is descriptive only. "
                f"{LOWER_LEVEL_NOTE}"
            ),
            imported_by_user_id=user.id,
            source_dataset=SOURCE_DATASET,
            source_file_name=extract.file_name,
            source_sha256=extract.sha256,
            source_sheet=extract.sheet,
        )
        session.add(version)
        session.flush()
        rows = []
        for item in report.units:
            for year in years:
                rows.append(
                    PopulationValue(
                        version_id=version.id,
                        org_unit_id=item.org_unit.id,
                        year=year,
                        population=item.unit.values[year],
                        source_unit_name=item.unit.name,
                        source_unit_type=item.unit.unit_type,
                        source_region=item.unit.region,
                        source_column_label=extract.column_labels[year],
                    )
                )
        if national_unit is not None:
            rows.extend(
                PopulationValue(
                    version_id=version.id,
                    org_unit_id=national_unit.id,
                    year=year,
                    population=extract.national_total[year],
                    source_unit_name=NATIONAL_TOTAL_LABEL,
                    source_unit_type="national_total",
                    source_region=None,
                    source_column_label=extract.column_labels[year],
                )
                for year in years
            )
        session.add_all(rows)
        session.flush()
        write_audit(
            session,
            actor_user_id=user.id,
            action="population_workbook_imported",
            resource_type="population_version",
            resource_id=str(version.id),
            after={
                "code": group_code,
                "status": ApprovalStatus.DRAFT.value,
                "source_sha256": extract.sha256,
                "source_file_name": extract.file_name,
                "values": len(rows),
                "national_total_included": national_unit is not None,
                "match_counts": report.counts,
            },
        )
        versions.append(version)
    return versions


# ---------------------------------------------------------------------------
# Structure validation, derived totals and governed staging
# ---------------------------------------------------------------------------


def region_totals(extract: WorkbookExtract) -> dict[str, dict[str, int]]:
    """Broad-region totals per year, derived from workbook membership only."""
    totals: dict[str, dict[str, int]] = {}
    for unit in extract.units:
        region = unit.region or "Unclassified"
        bucket = totals.setdefault(region, {})
        for year, value in unit.values.items():
            bucket[str(year)] = bucket.get(str(year), 0) + value
    return dict(sorted(totals.items()))


def derived_national_totals(extract: WorkbookExtract) -> dict[str, int]:
    """National population derived as the sum of the district and city rows."""
    totals: dict[str, int] = {}
    for unit in extract.units:
        for year, value in unit.values.items():
            totals[str(year)] = totals.get(str(year), 0) + value
    return dict(sorted(totals.items()))


def structure_findings(extract: WorkbookExtract) -> list[str]:
    """Differences between the workbook and its expected structure. Reported, never corrected."""
    findings: list[str] = []
    types = Counter(unit.unit_type for unit in extract.units)
    if types.get("District", 0) != EXPECTED_DISTRICTS:
        findings.append(f"Expected {EXPECTED_DISTRICTS} districts; found {types.get('District', 0)}.")
    if types.get("City", 0) != EXPECTED_CITIES:
        findings.append(f"Expected {EXPECTED_CITIES} cities; found {types.get('City', 0)}.")
    if len(extract.units) != EXPECTED_UNITS:
        findings.append(f"Expected {EXPECTED_UNITS} administrative units; found {len(extract.units)}.")
    if extract.years != EXPECTED_YEARS:
        findings.append(f"Expected years {EXPECTED_YEARS}; found {extract.years}.")
    unknown_regions = sorted(
        {unit.region for unit in extract.units if unit.region not in BROAD_REGIONS and unit.region is not None}
    )
    if unknown_regions:
        findings.append(f"Unexpected region labels: {', '.join(unknown_regions)}.")
    derived = derived_national_totals(extract)
    for year in extract.years:
        stated = extract.national_total.get(year)
        if stated is not None and derived.get(str(year)) != stated:
            findings.append(f"NATIONAL TOTAL for {year} does not equal the sum of the unit rows.")
    return findings


def detect_reference_scope(session: Session) -> str:
    """Whether candidate matches are against synthetic fixtures or a full hierarchy.

    A synthetic development database has a handful of organisation units; an authoritative
    hierarchy has the full district and city cohort. Synthetic matches are never reported as
    production matches.
    """
    candidates = session.scalars(
        select(OrgUnit).where(
            OrgUnit.active.is_(True),
            OrgUnit.level_type.in_([OrgUnitLevel.DISTRICT.value, OrgUnitLevel.CITY.value]),
        )
    ).all()
    return REFERENCE_AUTHORITATIVE if len(list(candidates)) >= EXPECTED_UNITS else REFERENCE_SYNTHETIC


def stage_population_workbook(
    session: Session,
    user: User,
    *,
    report: ReconciliationReport,
    source_display_name: str | None = None,
    notes: str | None = None,
) -> PopulationImportBatch:
    """Record every source row and its match state without creating any denominator.

    Staging is deliberately separate from applying: it preserves all 146 rows and the state of
    each crosswalk decision while the authoritative organisation-unit hierarchy is still absent.
    """
    require_action(session, user, ActionPermission.EDIT_POPULATION)
    extract = report.extract
    scope = detect_reference_scope(session)
    types = Counter(unit.unit_type for unit in extract.units)
    batch = PopulationImportBatch(
        source_dataset=SOURCE_DATASET,
        source_file_name=extract.file_name,
        source_display_name=source_display_name or extract.file_name,
        source_sha256=extract.sha256,
        source_type=SOURCE_TYPE,
        source_sheet=extract.sheet,
        importer_version=IMPORTER_VERSION,
        imported_at=datetime.now(UTC),
        imported_by_user_id=user.id,
        year_min=min(extract.years) if extract.years else None,
        year_max=max(extract.years) if extract.years else None,
        unit_count=len(extract.units),
        district_count=types.get("District", 0),
        city_count=types.get("City", 0),
        national_totals=derived_national_totals(extract),
        region_totals=region_totals(extract),
        match_counts=report.counts,
        reference_scope=scope,
        status=STAGED,
        review_status=REVIEW_PENDING,
        notes=notes,
    )
    session.add(batch)
    session.flush()
    findings = structure_findings(extract)
    for item in report.units:
        for year, value in sorted(item.unit.values.items()):
            session.add(
                PopulationImportRow(
                    batch_id=batch.id,
                    row_number=item.unit.row_number,
                    source_unit_name=item.unit.name,
                    source_unit_type=item.unit.unit_type,
                    source_region=item.unit.region,
                    year=year,
                    population=value,
                    source_column_label=extract.column_labels.get(year),
                    match_state=item.status,
                    candidate_org_unit_id=item.org_unit.id if item.org_unit else None,
                    alias_id=item.alias.id if item.alias else None,
                    review_state=REVIEW_PENDING,
                    validation_error=item.detail[:255] if item.detail else None,
                )
            )
    session.flush()
    write_audit(
        session,
        actor_user_id=user.id,
        action="population_workbook_staged",
        resource_type="population_import_batch",
        resource_id=str(batch.id),
        after={
            "source_file_name": extract.file_name,
            "source_sha256": extract.sha256,
            "units": len(extract.units),
            "cells": extract.cell_count,
            "reference_scope": scope,
            "match_counts": report.counts,
            "structure_findings": findings,
        },
        reason="Governed staging of an approved population source; no denominator was created.",
    )
    return batch


def staging_summary(session: Session, batch: PopulationImportBatch) -> dict:
    """Operator-facing summary of one staged batch."""
    rows = session.scalars(select(PopulationImportRow).where(PopulationImportRow.batch_id == batch.id)).all()
    states = Counter(row.match_state for row in rows)
    return {
        "batch_id": str(batch.id),
        "source_file_name": batch.source_file_name,
        "source_display_name": batch.source_display_name,
        "source_sha256": batch.source_sha256,
        "importer_version": batch.importer_version,
        "imported_at": batch.imported_at.isoformat() if batch.imported_at else None,
        "reference_scope": batch.reference_scope,
        "status": batch.status,
        "review_status": batch.review_status,
        "units": batch.unit_count,
        "districts": batch.district_count,
        "cities": batch.city_count,
        "years": [batch.year_min, batch.year_max],
        "staged_cells": len(rows),
        "match_states": dict(sorted(states.items())),
        "national_totals": batch.national_totals,
        "region_totals": batch.region_totals,
    }
