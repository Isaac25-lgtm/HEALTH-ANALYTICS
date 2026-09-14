from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from pptx import Presentation
from pptx.util import Inches, Pt
from sqlalchemy.orm import Session

from app.config import get_settings
from app.domain.enums import JobStatus
from app.models import ExportJob
from app.services.authorization import AuthorizationError
from app.version import SOFTWARE_VERSION

TEMPLATE_VERSION = "hpip-publish-1"

FILLS = {
    "green": PatternFill("solid", fgColor="C6EFCE"),
    "yellow": PatternFill("solid", fgColor="FFEB9C"),
    "red": PatternFill("solid", fgColor="FFC7CE"),
    "blue": PatternFill("solid", fgColor="BDD7EE"),
    "n_a": PatternFill("solid", fgColor="F2F2F2"),
}


def export_output_dir() -> Path:
    path = Path(get_settings().export_dir)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _scorecard_fill(row: dict) -> PatternFill | None:
    unit = str(row.get("unit") or "").casefold()
    if unit in {"count", "counts", "number"} or row.get("denominator_type") == "count":
        return FILLS["n_a"]
    return FILLS.get(row.get("quality_status") or row.get("status") or "n_a")


NOT_RECORDED = "Not recorded in the governed evidence"
PPTX_ROWS_PER_SLIDE = 8
MODULE_TITLES = {
    "anc": "ANC",
    "intrapartum": "Intrapartum and newborn",
    "immunization": "Immunization",
    "mpdsr": "MPDSR",
}


def _dashboard_rows(dashboard: dict) -> dict[str, dict]:
    rows = (dashboard.get("module_result") or {}).get("indicators") or []
    return {row.get("indicator_code"): row for row in rows if row.get("indicator_code")}


def _governed(row: dict, fallback: dict, key: str) -> object:
    """Governed evidence first, then the committed snapshot row; never a silent blank."""
    value = row.get(key)
    if value in (None, ""):
        value = fallback.get(key)
    if isinstance(value, dict):
        value = value.get("text") or value.get("state")
    return value if value not in (None, "") else NOT_RECORDED


def _freshness_text(package: dict, dashboard: dict, row: dict, fallback: dict) -> str:
    for candidate in (row.get("source_freshness_at"), fallback.get("source_freshness_at")):
        if candidate:
            return str(candidate)
    freshness = package.get("freshness") or (dashboard.get("module_result") or {}).get("freshness") or {}
    return str(freshness.get("latest_source_freshness_at") or freshness.get("latest_extracted_at") or NOT_RECORDED)


def _write_excel(_dashboard: dict, package: dict, destination: Path) -> None:
    book = Workbook()
    score = book.active
    score.title = "Scorecard"
    score.append(["Indicator", "Value", "Unit", "Status", "Numerator", "Denominator"])
    snapshot_rows = _dashboard_rows(_dashboard)
    for row in package["indicators"]:
        score.append(
            [
                row.get("name"),
                row.get("raw_value"),
                row.get("unit"),
                row.get("status"),
                row.get("numerator"),
                row.get("denominator"),
            ]
        )
        fill = _scorecard_fill(row)
        if fill:
            for cell in score[score.max_row]:
                cell.fill = fill
    calc = book.create_sheet("Calculations")
    calc.append(
        [
            "Code",
            "Display",
            "Exact value",
            "Numerator",
            "Denominator",
            "Formula version",
            "Thresholds",
            "Status rule",
            "Change interpretation",
            "Run",
            "BLUE / unavailable reason",
            "Reason code",
            "Formula version rule",
        ]
    )
    for row in package["indicators"]:
        fallback = snapshot_rows.get(row.get("indicator_code"), {})
        calc.append(
            [
                row.get("indicator_code"),
                row.get("display_value"),
                row.get("raw_value"),
                row.get("numerator"),
                row.get("denominator"),
                _governed(row, fallback, "formula_version"),
                _governed(row, fallback, "thresholds"),
                row.get("status"),
                (row.get("change") or {}).get("interpretation") or row.get("interpretation") or "not_interpreted",
                row.get("calculation_run_id"),
                row.get("blue_reason"),
                row.get("reason_code"),
                _governed(row, fallback, "formula_version_rule"),
            ]
        )
    raw = book.create_sheet("Raw Data")
    raw.append(
        [
            "Code",
            "Numerator",
            "Denominator",
            "Mapping version",
            "Source freshness",
            "Aggregation scope",
            "Identifiers",
            "Missing is not converted to zero",
        ]
    )
    for row in package["indicators"]:
        fallback = snapshot_rows.get(row.get("indicator_code"), {})
        raw.append(
            [
                row.get("indicator_code"),
                row.get("numerator"),
                row.get("denominator"),
                _governed(row, fallback, "mapping_version"),
                _freshness_text(package, _dashboard, row, fallback),
                _governed(row, fallback, "aggregation_policy"),
                "redacted",
                "retained",
            ]
        )
    pop = book.create_sheet("Population")
    population = package.get("population") or {}
    pop.append(
        [
            "Year",
            "Value",
            "Source",
            "Status",
            "Reason",
            "Version code",
            "Version ID",
            "Entry ID",
            "Type",
            "Approval",
        ]
    )
    pop.append(
        [
            population.get("year"),
            population.get("population"),
            population.get("source"),
            population.get("status"),
            population.get("reason"),
            population.get("version_code"),
            population.get("version_id") or population.get("population_version_id"),
            population.get("used_facility_entry_id") or population.get("facility_population_entry_id"),
            population.get("population_type"),
            population.get("approval_status"),
        ]
    )
    method = book.create_sheet("Methodology")
    method.append(["Template", TEMPLATE_VERSION])
    method.append(["Software", SOFTWARE_VERSION])
    method.append(["Note", "Formulas remain in the versioned indicator catalogue. AI did not calculate these values."])
    quality = book.create_sheet("Data Quality")
    quality.append(["Rule", "Severity", "Explanation", "Occurrences"])
    for flag in package.get("quality_flags") or []:
        quality.append(
            [flag.get("rule_id"), flag.get("severity"), flag.get("explanation"), flag.get("occurrences") or 1]
        )
    meta = book.create_sheet("Metadata")
    meta.append(["Field", "Value"])
    for key in ("period", "comparison_period", "module", "current_run_id", "analysis_snapshot_id", "view_hash"):
        meta.append([key, package.get(key) or _dashboard.get(key)])
    meta.append(["scope", package.get("scope", {}).get("org_unit_name")])
    meta.append(["indicator_count", len(package["indicators"])])
    meta.append(["generated_at", datetime.now(UTC).isoformat()])
    book.save(destination)


def _chunks(items: list, size: int) -> list[list]:
    return [items[index : index + size] for index in range(0, len(items), size)] or [[]]


def _write_pptx(_dashboard: dict, package: dict, destination: Path) -> None:
    deck = Presentation()
    module_title = MODULE_TITLES.get(str(package.get("module")), str(package.get("module")))
    title = deck.slides.add_slide(deck.slide_layouts[0])
    title.shapes.title.text = f"{package['scope'].get('org_unit_name')} {module_title} briefing"
    title.placeholders[1].text = (
        f"{package.get('period')} · module {package.get('module')} · "
        f"run {package.get('current_run_id')} · {TEMPLATE_VERSION}"
    )
    indicators = list(package["indicators"])
    pages = _chunks(indicators, PPTX_ROWS_PER_SLIDE)
    for page_number, page in enumerate(pages, start=1):
        table_slide = deck.slides.add_slide(deck.slide_layouts[5])
        suffix = f" ({page_number} of {len(pages)})" if len(pages) > 1 else ""
        table_slide.shapes.title.text = f"Verified scorecard{suffix}"
        table = table_slide.shapes.add_table(
            len(page) + 1, 4, Inches(0.5), Inches(1.5), Inches(9), Inches(0.45 * (len(page) + 1))
        ).table
        table.cell(0, 0).text = "Indicator"
        table.cell(0, 1).text = "Value"
        table.cell(0, 2).text = "Unit"
        table.cell(0, 3).text = "Status"
        for index, row in enumerate(page, start=1):
            table.cell(index, 0).text = str(row.get("name") or row.get("indicator_code") or "")
            table.cell(index, 1).text = str(row.get("display_value") or "No data")
            table.cell(index, 2).text = str(row.get("unit") or "")
            status_text = str(row.get("status") or "n_a")
            if row.get("reason_code"):
                status_text = f"{status_text} ({row['reason_code']})"
            table.cell(index, 3).text = status_text
    notes = deck.slides.add_slide(deck.slide_layouts[1])
    notes.shapes.title.text = "Evidence notes"
    notes.placeholders[1].text = (
        f"All {len(indicators)} indicators in the snapshot are included across {len(pages)} scorecard slide(s). "
        "AI may only fill designated narrative fields. Tables are populated from the calculation run. "
        "Official MoH slide templates are not yet supplied; this is the platform-default family."
    )
    for shape in notes.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                paragraph.font.size = Pt(16)
    deck.save(destination)


def _write_report(_dashboard: dict, package: dict, destination: Path) -> None:
    module_title = MODULE_TITLES.get(str(package.get("module")), str(package.get("module")))
    population = package.get("population") or {}
    lines = [
        f"# {module_title} narrative report — {package['scope'].get('org_unit_name')}",
        "",
        "Format: Markdown (.md). This is not a Word or PDF document.",
        "",
        f"Period: {package.get('period')}",
        f"Comparison: {package.get('comparison_period')}",
        f"Module: {package.get('module')}",
        f"Calculation run: {package.get('current_run_id')}",
        f"Software: {SOFTWARE_VERSION}",
        f"Template: {TEMPLATE_VERSION}",
        "",
        "## Executive summary",
        "This report uses only verified calculation-run values. AI did not calculate official numbers.",
        f"Snapshot: {package.get('analysis_snapshot_id') or _dashboard.get('analysis_snapshot_id')}",
        f"View hash: {package.get('view_hash') or _dashboard.get('view_hash')}",
        "",
        "## Key performance results",
    ]
    for row in package["indicators"]:
        reason = f"; {row.get('reason_code')}" if row.get("reason_code") else ""
        lines.append(
            f"- {row.get('name') or row.get('indicator_code')}: {row.get('display_value') or 'No data'} "
            f"{row.get('unit') or ''} ({row.get('status')}{reason})"
        )
    lines.extend(
        [
            "",
            "## Population",
            f"- Status: {population.get('status')}",
            f"- Year: {population.get('year')}",
            f"- Value: {population.get('population')}",
            f"- Source: {population.get('source')}",
            f"- Version: {population.get('version_code')}",
            f"- Reason: {population.get('reason')}",
            "",
            "## Data quality",
        ]
    )
    for flag in package.get("quality_flags") or []:
        lines.append(
            f"- {flag.get('rule_id')} ({flag.get('occurrences') or 1} occurrence(s)): {flag.get('explanation')}"
        )
    if not package.get("quality_flags"):
        lines.append("- No quality flags were attached to this run.")
    lines.extend(
        [
            "",
            "## Comparisons and contributors",
            "Child comparisons and contribution rows come from the persisted snapshot.",
            "No causal inference is added.",
            "",
            "## Methodology",
            f"Formula catalogue and template {TEMPLATE_VERSION}.",
            "Population and mapping versions are those stored on the snapshot.",
            "",
            "## Limitations",
            "Sensitive MPDSR identifiers and narratives are excluded. Official branding templates remain pending.",
            "",
            "## Recommendations",
            "Only verified snapshot findings are restated. None are invented.",
        ]
    )
    destination.write_text("\n".join(lines), encoding="utf-8")


ARTIFACT_WRITERS = {
    "excel": _write_excel,
    "powerpoint": _write_pptx,
    "report": _write_report,
}


def write_artifact(export_type: str, dashboard: dict, package: dict, destination: Path) -> None:
    """Write one export from committed snapshot content. Never recalculates."""
    writer = ARTIFACT_WRITERS.get(export_type)
    if writer is None:
        raise ValueError("Unsupported export type.")
    writer(dashboard, package, destination)


def generate_export_artifact(session: Session, job_id: UUID) -> ExportJob:
    """Run one claimed generation attempt in ``session`` (worker semantics, no automatic retry).

    Kept for callers that process a queued job synchronously. The attempt uses the same claim,
    temporary-file and durable-failure path as the Celery worker.
    """
    from app.services.export_jobs import ExportGenerationError, process_export_job

    job = session.get(ExportJob, job_id)
    if job is None:
        raise AuthorizationError("not_found", "Export job was not found.")
    if job.status == JobStatus.SUCCEEDED.value and job.file_path and Path(job.file_path).exists():
        return job
    attempt = process_export_job(job_id, session=session, auto_retry=False)
    session.expire_all()
    job = session.get(ExportJob, job_id)
    if attempt.outcome != "succeeded" and (job is None or job.status != JobStatus.SUCCEEDED.value):
        raise ExportGenerationError(attempt.error_code)
    return job
