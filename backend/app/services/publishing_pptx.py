"""A 16:9, colour-coded PowerPoint briefing written from a committed snapshot.

Everything is native and editable: tables, charts and the district map (vector freeform shapes).
No picture is embedded and no value is recalculated.
"""

from __future__ import annotations

import math
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from app.services.export_view import (
    INK,
    MUTED,
    NAVY,
    STATUS_FILL,
    STATUS_INK,
    STATUS_SOLID,
    ExportView,
    change_text,
    format_measure,
    interpretation_text,
    resolve_status,
    status_label,
    trend_series,
)

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
SCORECARD_ROWS_PER_SLIDE = 10
UNIT_ROWS_PER_SLIDE = 16
SCORECARD_TABLE_NAME = "Scorecard table"
UNITS_TABLE_NAME = "Units table"


def _rgb(hex_code: str) -> RGBColor:
    return RGBColor.from_string(hex_code)


def _chunks(items: list, size: int) -> list[list]:
    return [items[index : index + size] for index in range(0, len(items), size)] or [[]]


def _text(slide, left, top, width, height, text, *, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Emu(0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    return box


def _header(deck, view: ExportView, title: str, subtitle: str | None = None):
    """A title-only slide with a navy band, the slide title and the period line."""
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, Inches(0.95))
    band.fill.solid()
    band.fill.fore_color.rgb = _rgb(NAVY)
    band.line.fill.background()
    slide.shapes._spTree.remove(band._element)
    slide.shapes._spTree.insert(2, band._element)
    heading = slide.shapes.title
    heading.left, heading.top, heading.width, heading.height = Inches(0.5), Inches(0.12), Inches(12.3), Inches(0.72)
    heading.text = title
    frame = heading.text_frame
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        for run in paragraph.runs:
            run.font.size = Pt(24)
            run.font.bold = True
            run.font.color.rgb = _rgb("FFFFFF")
    _text(
        slide,
        Inches(0.5),
        Inches(1.02),
        Inches(12.3),
        Inches(0.35),
        subtitle or f"{view.scope_name} · {view.module_title} · {view.subtitle}",
        size=12,
        color=MUTED,
    )
    return slide


def _fill_cell(cell, text: str, *, status: str | None = None, bold=False, size=11, align=PP_ALIGN.LEFT, header=False):
    cell.text = text
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = align
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.bold = bold or header
        if header:
            run.font.color.rgb = _rgb("FFFFFF")
        elif status:
            run.font.color.rgb = _rgb(STATUS_INK.get(status, INK))
        else:
            run.font.color.rgb = _rgb(INK)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.06)
    cell.margin_top = cell.margin_bottom = Inches(0.02)
    cell.fill.solid()
    if header:
        cell.fill.fore_color.rgb = _rgb(NAVY)
    elif status:
        cell.fill.fore_color.rgb = _rgb(STATUS_FILL.get(status, STATUS_FILL["n_a"]))
    else:
        cell.fill.fore_color.rgb = _rgb("FFFFFF")


def _legend(slide, left, top):
    for index, status in enumerate(("green", "yellow", "red", "blue", "missing")):
        chip = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(1.55) * index, top, Inches(1.45), Inches(0.3)
        )
        chip.fill.solid()
        chip.fill.fore_color.rgb = _rgb(STATUS_FILL[status])
        chip.line.fill.background()
        frame = chip.text_frame
        frame.margin_top = frame.margin_bottom = Emu(0)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = status_label(status)
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = _rgb(STATUS_INK[status])


def _title_slide(deck, view: ExportView) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
    background.fill.solid()
    background.fill.fore_color.rgb = _rgb(NAVY)
    background.line.fill.background()
    slide.shapes._spTree.remove(background._element)
    slide.shapes._spTree.insert(2, background._element)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(4.35), Inches(2.2), Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = _rgb("F2C94C")
    accent.line.fill.background()
    title = slide.shapes.title
    title.left, title.top, title.width, title.height = Inches(0.8), Inches(2.2), Inches(11.7), Inches(2.0)
    title.text = view.title
    for paragraph in title.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        for run in paragraph.runs:
            run.font.size = Pt(40)
            run.font.bold = True
            run.font.color.rgb = _rgb("FFFFFF")
    subtitle = slide.placeholders[1]
    subtitle.left, subtitle.top, subtitle.width, subtitle.height = Inches(0.8), Inches(4.6), Inches(11.7), Inches(1.6)
    subtitle.text = (
        f"{view.subtitle}\nGenerated {view.generated_at} · calculation run {view.run_id}\n"
        "Ministry of Health · Health Performance Intelligence"
    )
    for paragraph in subtitle.text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        for run in paragraph.runs:
            run.font.size = Pt(16)
            run.font.color.rgb = _rgb("DCE8F5")


def _kpi_slide(deck, view: ExportView) -> None:
    slide = _header(deck, view, "Headline indicators")
    width, height, gap = Inches(3.95), Inches(2.45), Inches(0.25)
    for index, indicator in enumerate(view.kpis[:6]):
        column, row = index % 3, index // 3
        left = Inches(0.5) + (width + gap) * column
        top = Inches(1.6) + (height + gap) * row
        status = resolve_status(indicator)
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = _rgb("FFFFFF")
        card.line.color.rgb = _rgb("D5DEE8")
        edge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.09))
        edge.fill.solid()
        edge.fill.fore_color.rgb = _rgb(STATUS_SOLID.get(status, STATUS_SOLID["n_a"]))
        edge.line.fill.background()
        _text(
            slide,
            left + Inches(0.2),
            top + Inches(0.2),
            width - Inches(0.4),
            Inches(0.4),
            str(indicator.get("name") or indicator.get("indicator_code")),
            size=14,
            bold=True,
            color=NAVY,
        )
        _text(
            slide,
            left + Inches(0.2),
            top + Inches(0.65),
            width - Inches(0.4),
            Inches(0.8),
            format_measure(indicator),
            size=36,
            bold=True,
            color=INK,
        )
        change = change_text(indicator)
        interpretation = interpretation_text(indicator)
        comparison = f"{change} · {interpretation}" if interpretation else change
        if view.comparison_label and change != "No comparison":
            comparison = f"{comparison} vs {view.comparison_label}"
        _text(
            slide,
            left + Inches(0.2),
            top + Inches(1.5),
            width - Inches(0.4),
            Inches(0.35),
            comparison,
            size=11,
            color=MUTED,
        )
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.2), top + Inches(1.9), Inches(1.9), Inches(0.34)
        )
        pill.fill.solid()
        pill.fill.fore_color.rgb = _rgb(STATUS_FILL.get(status, STATUS_FILL["n_a"]))
        pill.line.fill.background()
        paragraph = pill.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = status_label(status)
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = _rgb(STATUS_INK.get(status, INK))


def _scorecard_slides(deck, view: ExportView) -> int:
    pages = _chunks(list(view.indicators), SCORECARD_ROWS_PER_SLIDE)
    for page_number, page in enumerate(pages, start=1):
        suffix = f" ({page_number} of {len(pages)})" if len(pages) > 1 else ""
        slide = _header(deck, view, f"Verified scorecard{suffix}")
        shape = slide.shapes.add_table(
            len(page) + 1, 5, Inches(0.5), Inches(1.5), Inches(12.3), Inches(0.42 * (len(page) + 1))
        )
        shape.name = SCORECARD_TABLE_NAME
        table = shape.table
        for index, width in enumerate((4.8, 1.6, 2.1, 1.9, 1.9)):
            table.columns[index].width = Inches(width)
        comparison_head = f"Comparison ({view.comparison_label})" if view.comparison_label else "Comparison"
        for column, label in enumerate(("Indicator", "Value", "Status", comparison_head, "Change")):
            _fill_cell(table.cell(0, column), label, header=True, align=PP_ALIGN.CENTER if column else PP_ALIGN.LEFT)
        for row_index, indicator in enumerate(page, start=1):
            status = resolve_status(indicator)
            comparison = indicator.get("comparison") or {}
            _fill_cell(
                table.cell(row_index, 0), str(indicator.get("name") or indicator.get("indicator_code") or ""), bold=True
            )
            _fill_cell(
                table.cell(row_index, 1), format_measure(indicator), status=status, bold=True, align=PP_ALIGN.CENTER
            )
            label = status_label(status)
            if indicator.get("reason_code") and status in {"missing", "blue", "n_a"}:
                label = f"{label} ({str(indicator['reason_code']).replace('_', ' ')})"
            _fill_cell(table.cell(row_index, 2), label, status=status, bold=True, size=10, align=PP_ALIGN.CENTER)
            comparison_status = resolve_status(comparison) if comparison.get("raw_value") is not None else None
            _fill_cell(
                table.cell(row_index, 3), format_measure(comparison), status=comparison_status, align=PP_ALIGN.CENTER
            )
            _fill_cell(table.cell(row_index, 4), change_text(indicator), align=PP_ALIGN.CENTER)
        _legend(slide, Inches(0.5), Inches(6.85))
    return len(pages)


def _project(features: list[dict]):
    """Equirectangular projection scaled by cos(latitude); fine at Uganda's latitude."""
    points = []

    def walk(node):
        if isinstance(node, list) and node and isinstance(node[0], (int, float)):
            points.append(node)
        elif isinstance(node, list):
            for item in node:
                walk(item)

    for feature in features:
        walk((feature.get("geometry") or {}).get("coordinates"))
    if not points:
        return None
    lons = [point[0] for point in points]
    lats = [point[1] for point in points]
    factor = math.cos(math.radians((min(lats) + max(lats)) / 2))
    return min(lons), max(lons), min(lats), max(lats), factor


def _rings(geometry: dict) -> list[list]:
    kind = (geometry or {}).get("type")
    coordinates = (geometry or {}).get("coordinates") or []
    if kind == "Polygon":
        return [coordinates[0]] if coordinates else []
    if kind == "MultiPolygon":
        return [polygon[0] for polygon in coordinates if polygon]
    return []


def _map_slide(deck, view: ExportView) -> None:
    if not view.map_features:
        return
    slide = _header(deck, view, f"Geographic performance · {view.map_indicator_name}")
    bounds = _project(view.map_features)
    if bounds is None:
        return
    min_lon, max_lon, min_lat, max_lat, factor = bounds
    box_left, box_top, box_width, box_height = Inches(0.5), Inches(1.5), Inches(7.2), Inches(5.2)
    span_x = (max_lon - min_lon) * factor or 1
    span_y = (max_lat - min_lat) or 1
    scale = min(box_width / span_x, box_height / span_y)
    offset_x = box_left + (box_width - span_x * scale) / 2
    offset_y = box_top + (box_height - span_y * scale) / 2

    def to_emu(point):
        x = offset_x + (point[0] - min_lon) * factor * scale
        y = offset_y + (max_lat - point[1]) * scale
        return int(x), int(y)

    for feature in view.map_features:
        properties = feature.get("properties") or {}
        status = resolve_status(properties)
        for ring in _rings(feature.get("geometry") or {}):
            vertices = [to_emu(point) for point in ring if len(point) >= 2]
            if len(vertices) < 3:
                continue
            builder = slide.shapes.build_freeform(vertices[0][0], vertices[0][1], scale=1.0)
            builder.add_line_segments(vertices[1:], close=True)
            shape = builder.convert_to_shape()
            shape.name = f"{properties.get('name')}: {format_measure(properties)}"
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(STATUS_FILL.get(status, STATUS_FILL["missing"]))
            shape.line.color.rgb = _rgb("FFFFFF")
            shape.line.width = Pt(0.5)
    _legend(slide, Inches(0.5), Inches(6.85))
    _performer_lists(slide, view, Inches(8.1), Inches(1.5))
    if view.map_note:
        _text(slide, Inches(8.1), Inches(6.45), Inches(4.8), Inches(0.3), view.map_note, size=9, color=MUTED)


def _performer_lists(slide, view: ExportView, left, top) -> None:
    ranking = view.ranking
    if not ranking.get("ranking_allowed"):
        _text(
            slide,
            left,
            top,
            Inches(4.8),
            Inches(0.6),
            ranking.get("reason") or "Ranking is not available for this indicator.",
            size=11,
            color=MUTED,
        )
        return
    y = top
    for title, rows in (
        ("Best performing", ranking.get("best") or []),
        ("Needs attention", ranking.get("worst") or []),
    ):
        _text(slide, left, y, Inches(4.8), Inches(0.35), title, size=13, bold=True, color=NAVY)
        y += Inches(0.4)
        for entry in rows[:5]:
            status = resolve_status(entry)
            _text(slide, left, y, Inches(3.3), Inches(0.3), str(entry.get("org_unit_name") or ""), size=11)
            chip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(3.4), y, Inches(1.35), Inches(0.28))
            chip.fill.solid()
            chip.fill.fore_color.rgb = _rgb(STATUS_FILL.get(status, STATUS_FILL["n_a"]))
            chip.line.fill.background()
            paragraph = chip.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            run = paragraph.add_run()
            run.text = format_measure(entry)
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = _rgb(STATUS_INK.get(status, INK))
            y += Inches(0.34)
        y += Inches(0.25)


def _unit_slides(deck, view: ExportView) -> None:
    if not view.units or not view.unit_columns:
        return
    pages = _chunks(view.units, UNIT_ROWS_PER_SLIDE)
    columns = view.unit_columns
    for page_number, page in enumerate(pages, start=1):
        suffix = f" ({page_number} of {len(pages)})" if len(pages) > 1 else ""
        slide = _header(deck, view, f"{view.unit_label} scorecard{suffix}")
        shape = slide.shapes.add_table(
            len(page) + 1, len(columns) + 1, Inches(0.5), Inches(1.45), Inches(12.3), Inches(0.31 * (len(page) + 1))
        )
        shape.name = UNITS_TABLE_NAME
        table = shape.table
        table.columns[0].width = Inches(3.3)
        for index in range(len(columns)):
            table.columns[index + 1].width = Inches(9.0 / len(columns))
        _fill_cell(table.cell(0, 0), "Unit", header=True, size=10)
        for index, column in enumerate(columns, start=1):
            _fill_cell(
                table.cell(0, index),
                str(column.get("name") or column.get("indicator_code")),
                header=True,
                size=10,
                align=PP_ALIGN.CENTER,
            )
        for row_index, unit in enumerate(page, start=1):
            _fill_cell(table.cell(row_index, 0), str(unit.get("org_unit_name") or ""), bold=True, size=10)
            values = unit.get("values") or {}
            for index, column in enumerate(columns, start=1):
                measure = values.get(column.get("indicator_code"))
                _fill_cell(
                    table.cell(row_index, index),
                    format_measure(measure),
                    status=resolve_status(measure),
                    bold=True,
                    size=10,
                    align=PP_ALIGN.CENTER,
                )
        _legend(slide, Inches(0.5), Inches(6.95))


def _bar_chart(slide, rows: list[dict], left, top, width, height, title: str, axis_max: float) -> None:
    usable = [row for row in rows if row.get("raw_value") is not None]
    if not usable:
        return
    data = CategoryChartData()
    data.categories = [str(row.get("org_unit_name") or row.get("org_unit_code") or "") for row in reversed(usable)]
    data.add_series("Value", [row.get("raw_value") for row in reversed(usable)])
    frame = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, left, top, width, height, data)
    chart = frame.chart
    chart.has_legend = False
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.color.rgb = _rgb(NAVY)
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(11)
    plot.data_labels.font.bold = True
    plot.data_labels.number_format = '0.0"%"' if usable[0].get("unit") == "%" else "#,##0"
    plot.data_labels.number_format_is_linked = False
    for index, row in enumerate(reversed(usable)):
        point = plot.series[0].points[index]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = _rgb(STATUS_SOLID.get(resolve_status(row), STATUS_SOLID["n_a"]))
    chart.category_axis.tick_labels.font.size = Pt(11)
    # Bars start at zero on a scale shared by both charts, so lengths compare honestly.
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = axis_max
    chart.value_axis.visible = False
    chart.value_axis.has_major_gridlines = False


def _ranking_slide(deck, view: ExportView) -> None:
    ranking = view.ranking
    if not ranking.get("ranking_allowed"):
        return
    slide = _header(deck, view, f"Top and bottom performers · {view.ranking_indicator_name}")
    best, worst = ranking.get("best") or [], ranking.get("worst") or []
    values = [row["raw_value"] for row in best + worst if row.get("raw_value") is not None]
    percent = any(row.get("unit") == "%" for row in best + worst)
    axis_max = max([100.0 if percent else 0.0, *values]) * 1.15 or 1.0
    _bar_chart(slide, best, Inches(0.5), Inches(1.5), Inches(6.0), Inches(5.2), "Best performing", axis_max)
    _bar_chart(slide, worst, Inches(6.8), Inches(1.5), Inches(6.0), Inches(5.2), "Needs attention", axis_max)


def _trend_slide(deck, view: ExportView) -> None:
    series = trend_series(view)
    if not series or all(value is None for _label, value in series):
        return
    indicator = view.trend_indicator or {}
    slide = _header(deck, view, f"Trend · {indicator.get('name') or indicator.get('indicator_code')}")
    data = CategoryChartData()
    data.categories = [label for label, _value in series]
    data.add_series(str(indicator.get("name") or "Value"), [value for _label, value in series])
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS, Inches(0.5), Inches(1.5), Inches(12.3), Inches(5.3), data
    )
    chart = frame.chart
    chart.has_legend = False
    line = chart.plots[0].series[0]
    line.format.line.color.rgb = _rgb(NAVY)
    line.format.line.width = Pt(2.5)
    line.smooth = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = _rgb("E2E8F0")
    chart.value_axis.tick_labels.font.size = Pt(11)
    chart.category_axis.tick_labels.font.size = Pt(11)
    if indicator.get("unit") == "%":
        chart.value_axis.tick_labels.number_format = '0"%"'
        chart.value_axis.tick_labels.number_format_is_linked = False
    _text(
        slide,
        Inches(0.5),
        Inches(6.9),
        Inches(12.3),
        Inches(0.3),
        "Missing periods are gaps, never zero or interpolated.",
        size=10,
        color=MUTED,
    )


def _notes_slide(deck, view: ExportView, pages: int, template_version: str) -> None:
    slide = _header(deck, view, "Evidence notes")
    lines = [
        f"All {len(view.indicators)} indicators in the snapshot are included across {pages} scorecard slide(s).",
        f"Calculation run {view.run_id}; snapshot {view.snapshot_id}.",
        "Tables, charts and the map are populated from the committed calculation run. Nothing was recalculated.",
        "Colours follow each value's approved status. AI may only fill designated narrative fields.",
        f"Official MoH slide templates are not yet supplied; this is the platform-default family ({template_version}).",
    ]
    population = view.population
    if population.get("population"):
        lines.append(
            f"Population {population.get('population'):,} ({population.get('year')}) · "
            f"{population.get('source') or 'source not stated'}."
        )
    box = slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(12.3), Inches(5.0))
    frame = box.text_frame
    frame.word_wrap = True
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(10)
        run = paragraph.add_run()
        run.text = f"• {line}"
        run.font.size = Pt(16)
        run.font.color.rgb = _rgb(INK)


def write_pptx(view: ExportView, destination: Path, template_version: str) -> None:
    deck = Presentation()
    deck.slide_width = SLIDE_WIDTH
    deck.slide_height = SLIDE_HEIGHT
    _title_slide(deck, view)
    _kpi_slide(deck, view)
    pages = _scorecard_slides(deck, view)
    _map_slide(deck, view)
    _ranking_slide(deck, view)
    _trend_slide(deck, view)
    _unit_slides(deck, view)
    _notes_slide(deck, view, pages, template_version)
    deck.save(destination)
